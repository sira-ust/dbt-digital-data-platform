"""Write UST_<MON>_Promo_Social_Listening_Match.xlsx from mart_promo_social_match (step 5).

Three sheets, in the layout of the reviewed October workbook (2026-10-07):
  1. "Summary"                    opens first. Counts per status, per promo sheet and
                                  per channel, all LIVE formulas over sheet 2, so rows
                                  added afterwards (TikTok) are counted without
                                  re-running anything; then the items nothing was found
                                  for, as of the export.
  2. "Promo vs Social Listening"  one row per matched post / video, or one Unmatched
                                  row per item. Views and Likes as numbers, the
                                  platform's own wording beside them, links as "Open ↗";
                                  every other item's rows shaded, so an item reads as
                                  one block; Unmatched rows greyed.
  3. "TikTok Search List"         the items to search on TikTok by hand with Claude in
                                  Chrome (prompt: scripts/prompts/tiktok_promo_search.md):
                                  every branded item social listening did not match,
                                  with search phrases, required words, exclusions and
                                  a TikTok search link. The prompt has Claude delete
                                  this sheet when done, leaving the 2-sheet deliverable.

Unlike the October workbook, "Promo Items" on the Summary is a formula too (a
distinct count of Item No. per status). October typed numbers in, which go stale
the moment a TikTok row turns an Unmatched item into Exact.

    python scripts/export_promo_social_match.py                       # latest month, local
    python scripts/export_promo_social_match.py --promo-month 2026-10-01 --out-dir C:/somewhere
    python scripts/export_promo_social_match.py --backend databricks  # -> the promo_output Volume
"""

from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
import urllib.parse
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path.cwd() / "scripts"))
try:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
except NameError:
    pass

from promo_common import (  # noqa: E402
    INTERMEDIATE_SCHEMA, REPORTING_SCHEMA, STAGING_SCHEMA, local_data_root, query,
    split_terms,
)

DBX_OUT_DIR = "/Volumes/ust_databricks/social/promo_output"
SHEET_MAIN, SHEET_SUMMARY, SHEET_TIKTOK = "Promo vs Social Listening", "Summary", "TikTok Search List"

# The layout of the reviewed October workbook (UST_OCT_..._Match 1_2.xlsx, 2026-10-07).
# Columns A-J keep their letters, so the Summary's formulas (A = Promo Sheet,
# B = Item No., F = Match Status, H = Channel) hold however many rows are added.
HEADERS = ["Promo Sheet", "Item No.", "{MON} Promo Item", "Size", "Promo Code", "Match Status",
           "Matched Content Item", "Channel", "Profile", "Post Date", "Views", "Likes",
           "Approx. Count", "Engagement (as shown)", "Caption", "Link", "Source", "Match Note"]
WIDTHS = [13, 10, 32, 15, 10, 12, 36, 14, 22, 11, 13, 10, 9, 22, 55, 11, 15, 48]
COL = {h: i for i, h in enumerate(HEADERS, start=1)}
# status -> (fill, font colour)
# ARGB, opaque, as Excel writes them (openpyxl would store a 6-digit RGB as 00RRGGBB)
STATUS_STYLE = {"Exact": ("FFC6EFCE", "FF006100"), "Close": ("FFFFEB9C", "FF9C5700"),
                "Unmatched": ("FFEDEDED", "FF595959")}
HEADER_FILL, BORDER_COLOR, LINK_COLOR = "FF1F4E78", "FFBFBFBF", "FF0563C1"
BAND_FILL = "FFEEF3F8"       # every other ITEM, all of its rows, so an item reads as one block
MUTED, NOTE_GREY = "FF808080", "FF595959"
LINK_TEXT = "Open ↗"
HEADER_HEIGHT, ROW_HEIGHT, SUMMARY_ROW_HEIGHT = 27.75, 39.75, 15
FORMULA_ROWS = 5000          # Summary formulas cover this many sheet-1 rows

TIKTOK_HEADERS = ["Item No.", "{MON} Promo Item", "Size", "Brand", "Product", "Flavor / Variant",
                  "Current Status", "Search Phrases", "Must Mention (brand)", "Flavor Words",
                  "Exclude If", "TikTok Search Link"]
TIKTOK_WIDTHS = [10, 34, 16, 14, 22, 22, 14, 40, 36, 36, 30, 45]


# ─────────────────────────────────────────────────────────────────────────────
# Read
# ─────────────────────────────────────────────────────────────────────────────

def _rel(backend, dbx_schema, local_schema, name):
    return f"{dbx_schema}.{name}" if backend == "databricks" else f"{local_schema}.{name}"


def read_inputs(backend: str, promo_month: str | None) -> dict:
    mart = _rel(backend, REPORTING_SCHEMA, "ust_reporting", "mart_promo_social_match")
    items_rel = _rel(backend, INTERMEDIATE_SCHEMA, "ust_intermediate", "int_promo_items_enriched")
    if not promo_month:
        promo_month = str(query(backend, f"select max(promo_month) as m from {mart}")[0]["m"])[:10]
    rows = query(backend, f"""
        select * from {mart}
        where promo_month = date '{promo_month}'
        order by sheet_order, row_order, match_rank
    """)
    if not rows:
        raise SystemExit(f"no rows in {mart} for {promo_month}")
    items = query(backend, f"""
        select * from {items_rel}
        where promo_month = date '{promo_month}'
        order by sheet_order, row_order
    """)
    mentions = query(backend, f"""
        select count(*) as n, min(posted_at) as first_at, max(posted_at) as last_at
        from {_rel(backend, STAGING_SCHEMA, 'ust_staging', 'stg_mentionlytics__mentions')}
    """)[0]
    yt_rel = _rel(backend, STAGING_SCHEMA, "ust_staging", "stg_promo__youtube_videos")
    searched = query(backend, f"select max(searched_at) as s from {yt_rel} "
                              f"where promo_month = date '{promo_month}'")[0]["s"]
    return {"promo_month": promo_month, "rows": rows, "items": items, "mentions": mentions,
            "searched_at": searched}


# ─────────────────────────────────────────────────────────────────────────────
# Write
# ─────────────────────────────────────────────────────────────────────────────

def month_label(promo_month: str) -> str:
    return date.fromisoformat(promo_month[:10]).strftime("%b").upper()


def _fmt_date(v) -> str:
    if v is None or str(v) in ("", "NaT", "None"):
        return "n/a"
    d = v if isinstance(v, (date, datetime)) else datetime.fromisoformat(str(v)[:19])
    return f"{d:%b} {d.day}, {d.year}"


def footnote(data: dict) -> str:
    items = data["items"]
    src = sorted({i["source_file"] for i in items if i.get("source_file")})
    m = data["mentions"]
    scraped = (f"Rows with Source = Scraped came from YouTube searches on {_fmt_date(data['searched_at'])}, "
               f"keeping up to 3 of the highest-view videos that name the exact product."
               if data["searched_at"] else "No YouTube search has run for this month yet.")
    return (f"Promo item counts are distinct Item Nos. Source: {', '.join(src) or 'promo workbook'} "
            f"({len({i['item_no'] for i in items})} items) vs social listening "
            f"({int(m['n'] or 0):,} posts, {_fmt_date(m['first_at'])} – {_fmt_date(m['last_at'])}). "
            f"{scraped} YouTube shows views; TikTok shows likes. TikTok rows are added by hand "
            f"from the TikTok Search List sheet.")


def _none(v):
    if v is None:
        return None
    try:
        import math
        if isinstance(v, float) and math.isnan(v):
            return None
    except Exception:
        pass
    s = str(v)
    return None if s in ("", "NaT", "nan", "None") else v


def _count(v):
    """A metric as an integer for the Views / Likes columns, or None when missing or 0."""
    v = _none(v)
    if v is None:
        return None
    try:
        n = int(float(v))
    except (TypeError, ValueError):
        return None
    return n if n > 0 else None


def build_workbook(data: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    mon = month_label(data["promo_month"])
    thin = Side(style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    body_font, bold_font = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    head_fill = PatternFill("solid", fgColor=HEADER_FILL)
    band_fill = PatternFill("solid", fgColor=BAND_FILL)
    wrap_top = Alignment(vertical="top", wrap_text=True)
    link_font = Font(name="Arial", size=10, color=LINK_COLOR, underline="single")

    def header_row(ws, headers, widths, row=1, align=None):
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=c, value=h.replace("{MON}", mon))
            cell.font, cell.fill, cell.border = head_font, head_fill, border
            cell.alignment = align or Alignment(vertical="center", wrap_text=True)
        for c, w in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(c)].width = w

    def body_cell(ws, r, c, value, align=wrap_top):
        cell = ws.cell(row=r, column=c, value=value)
        cell.font, cell.alignment, cell.border = body_font, align, border
        return cell

    wb = Workbook()
    sm = wb.active                          # Summary is the FIRST sheet, and opens first
    sm.title = SHEET_SUMMARY
    ws = wb.create_sheet(SHEET_MAIN)

    # ── Promo vs Social Listening ────────────────────────────────────────────
    header_row(ws, HEADERS, WIDTHS)
    ws.row_dimensions[1].height = HEADER_HEIGHT
    numbers = Alignment(horizontal="right", vertical="top")
    centred = Alignment(horizontal="center", vertical="top")
    band, last_item = False, None
    for r, row in enumerate(data["rows"], start=2):
        key = (row["promo_sheet"], row["row_order"])
        if key != last_item:                 # a new item: flip the band
            band, last_item = (not band) if last_item is not None else False, key
        unmatched = row["match_status"] == "Unmatched"
        views, likes = (None, None) if unmatched else (_count(row.get("views")), _count(row.get("likes")))
        if not unmatched and views is None and row.get("source") == "Scraped" \
                and str(row.get("channel") or "").startswith("YouTube"):
            views = 0                        # YouTube always reports views, even none
        values = {
            "Promo Sheet": row["promo_sheet"],
            "Item No.": int(row["item_no_int"]) if _none(row.get("item_no_int")) is not None else row["item_no"],
            "{MON} Promo Item": (row["promo_description"] or "").strip(),
            "Size": _none(row["promo_size"]),
            "Promo Code": _none(row["promo_code"]),
            "Match Status": row["match_status"],
        }
        if not unmatched:
            values.update({
                "Matched Content Item": _none(row["matched_content_item"]),
                "Channel": _none(row["channel"]),
                "Profile": _none(row["profile"]),
                "Post Date": _none(row["post_date"]),
                "Views": views,
                "Likes": likes,
                "Approx. Count": None,       # set by the TikTok step when TikTok shows only "~3.7K"
                "Engagement (as shown)": _none(row["engagement_text"]),
                "Caption": _none(row["caption"]),
                "Link": LINK_TEXT if _none(row["link"]) else None,
                "Source": _none(row["source"]),
                "Match Note": _none(row["match_note"]),
            })
        for h, c in COL.items():
            align = numbers if h in ("Views", "Likes") else centred if h == "Approx. Count" else wrap_top
            cell = body_cell(ws, r, c, values.get(h), align)
            if band:
                cell.fill = band_fill
            if unmatched:
                cell.font = Font(name="Arial", size=10, color=MUTED)
            if h in ("Views", "Likes"):
                cell.number_format = "#,##0"
            elif h == "Match Status":
                fill, colour = STATUS_STYLE.get(row["match_status"], ("FFFFFFFF", "FF000000"))
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(name="Arial", size=10, bold=True, color=colour)
            elif h == "Link" and cell.value:
                cell.hyperlink, cell.font = str(row["link"]), link_font
        ws.row_dimensions[r].height = ROW_HEIGHT
    last = len(data["rows"]) + 1
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last}"

    # ── Summary ──────────────────────────────────────────────────────────────
    # Every count is a LIVE formula over the main sheet, so rows the TikTok step adds
    # are counted without re-running anything.
    q = f"'{SHEET_MAIN}'"
    rng_a, rng_b, rng_f = (f"{q}!${c}$2:${c}${FORMULA_ROWS}" for c in "ABF")
    distinct = f'COUNTIFS({rng_b},{rng_b}&"",{rng_f},{rng_f}&"")'   # rows per (item, status)
    plain = Alignment(vertical="bottom")
    statuses = ["Exact", "Close", "Unmatched"]

    def title(r, text):
        sm.cell(row=r, column=1, value=text).font = Font(name="Arial", size=11, bold=True)

    def total_row(r, first, last_r, cols):
        body_cell(sm, r, 1, "Total", plain).font = bold_font
        for c in range(2, cols + 1):
            col = get_column_letter(c)
            body_cell(sm, r, c, f"=SUM({col}{first}:{col}{last_r})", plain)

    header_row(sm, ["Match Status", "Promo Items", "Matched Posts (rows)"], [16, 34, 22, 10],
               align=Alignment(horizontal="center", vertical="center", wrap_text=True))
    for r, status in enumerate(statuses, start=2):
        body_cell(sm, r, 1, status)
        body_cell(sm, r, 2, f"=SUMPRODUCT(({rng_f}=A{r})/{distinct})")      # distinct Item Nos.
        body_cell(sm, r, 3, f'=IF(A{r}="Unmatched",0,COUNTIF({q}!F:F,A{r}))')
    body_cell(sm, 5, 1, "Total").font = bold_font
    body_cell(sm, 5, 2, "=SUM(B2:B4)")
    body_cell(sm, 5, 3, "=SUM(C2:C4)")
    note = sm.cell(row=7, column=1, value=footnote(data))
    note.font = Font(name="Arial", size=10, italic=True)
    note.alignment = Alignment(vertical="top", wrap_text=False)

    # promo items by sheet, one column per promo sheet in workbook order
    sheets = list(dict.fromkeys(row["promo_sheet"] for row in data["rows"]))
    title(9, "Promo items by sheet")
    header_row(sm, ["Match Status"] + sheets + ["Total"], None, row=10, align=plain)
    for i, status in enumerate(statuses):
        r = 11 + i
        body_cell(sm, r, 1, status, plain)
        for c, sheet in enumerate(sheets, start=2):
            body_cell(sm, r, c, f'=SUMPRODUCT(({rng_a}="{sheet}")*({rng_f}=$A{r})/{distinct})', plain)
        body_cell(sm, r, len(sheets) + 2,
                  f"=SUM(B{r}:{get_column_letter(len(sheets) + 1)}{r})", plain)
    total_row(14, 11, 13, len(sheets) + 2)

    # matched posts by channel; TikTok is always listed, because the TikTok step adds it
    channels = sorted({str(row["channel"]) for row in data["rows"]
                       if row["match_status"] != "Unmatched" and _none(row.get("channel"))} | {"TikTok"},
                      key=str.casefold)
    title(16, "Matched posts by channel")
    header_row(sm, ["Channel", "Exact", "Close", "Total"], None, row=17, align=plain)
    for i, channel in enumerate(channels):
        r = 18 + i
        body_cell(sm, r, 1, channel, plain)
        for c, status in ((2, "Exact"), (3, "Close")):
            body_cell(sm, r, c, f'=COUNTIFS({q}!$H:$H,$A{r},{q}!$F:$F,"{status}")', plain)
        body_cell(sm, r, 4, f"=SUM(B{r}:C{r})", plain)
    end = 18 + len(channels)
    total_row(end, 18, end - 1, 4)

    # the items nothing was found for, as of this export (the TikTok step updates it)
    start = end + 2
    title(start, "Unmatched promo items (no post found)")
    header_row(sm, ["Item No.", f"{mon} Promo Item", "Promo Sheet"], None, row=start + 1, align=plain)
    r = start + 2
    for row in data["rows"]:
        if row["match_status"] == "Unmatched":
            item_no = int(row["item_no_int"]) if _none(row.get("item_no_int")) is not None else row["item_no"]
            body_cell(sm, r, 1, item_no, plain)
            body_cell(sm, r, 2, (row["promo_description"] or "").strip(), plain)
            body_cell(sm, r, 3, row["promo_sheet"], plain)
            r += 1
    tail = sm.cell(row=r, column=1, value="List captured when the workbook was built; to see current gaps, "
                                          f"filter Match Status = Unmatched on the {SHEET_MAIN} tab.")
    tail.font = Font(name="Arial", size=9, italic=True, color=NOTE_GREY)
    for i in range(1, r + 1):
        if any(sm.cell(row=i, column=c).value is not None for c in range(1, 5)):
            sm.row_dimensions[i].height = SUMMARY_ROW_HEIGHT

    # ── TikTok Search List (deleted by the TikTok step when it is done) ─────
    tk = wb.create_sheet(SHEET_TIKTOK)
    header_row(tk, TIKTOK_HEADERS, TIKTOK_WIDTHS,
               align=Alignment(horizontal="center", vertical="center", wrap_text=True))
    status_by_row = {}
    for row in data["rows"]:
        key = (row["promo_sheet"], row["row_order"])
        status_by_row.setdefault(key, (row["match_status"], row["source"]))
    r = 2
    for it in data["items"]:
        status, source = status_by_row.get((it["promo_sheet"], it["row_order"]), ("Unmatched", None))
        if not it.get("is_branded") or source == "Social Listening":
            continue                         # unbranded, or social listening already matched it
        phrases = split_terms(it.get("search_queries"))
        current = "Unmatched" if status == "Unmatched" else f"{status} (YouTube)"
        values = [
            int(it["item_no"]) if str(it["item_no"]).isdigit() else it["item_no"],
            (it["promo_description"] or "").strip(), _none(it.get("promo_size")),
            _none(it.get("brand")), _none(it.get("product")), _none(it.get("flavor_variant")),
            current, "\n".join(phrases),
            ", ".join(split_terms(it.get("brand_terms"))),
            ", ".join(split_terms(it.get("flavor_terms"))) or "(no flavor to require)",
            ", ".join(split_terms(it.get("negative_terms"))) or "-",
            f"https://www.tiktok.com/search/video?q={urllib.parse.quote(phrases[0])}" if phrases else None,
        ]
        for c, v in enumerate(values, start=1):
            cell = body_cell(tk, r, c, v)
            if c == 12 and v:
                cell.hyperlink, cell.font = v, link_font
        r += 1
    tk.freeze_panes = "C2"
    tk.auto_filter.ref = f"A1:L{max(r - 1, 1)}"
    return wb


def output_name(promo_month: str) -> str:
    return f"UST_{month_label(promo_month)}_Promo_Social_Listening_Match.xlsx"


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--backend", choices=["local", "databricks"], default="local")
    ap.add_argument("--promo-month", help="YYYY-MM-01 (default: the latest month in the mart)")
    ap.add_argument("--out-dir", help="folder to write to (default: <local_data_root>/promo/output, "
                                      f"or {DBX_OUT_DIR} on Databricks)")
    args = ap.parse_args(argv)

    data = read_inputs(args.backend, args.promo_month)
    out_dir = Path(args.out_dir) if args.out_dir else (
        Path(DBX_OUT_DIR) if args.backend == "databricks" else local_data_root() / "promo" / "output")
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / output_name(data["promo_month"])
    # An .xlsx is a zip, and writing one seeks back to patch headers. A Unity Catalog
    # Volume is write-once, sequential storage: saving straight to it fails with
    # "[Errno 5] Input/output error" (first Databricks run, 2026-10-07). So the workbook
    # is built on local disk and then copied over in one sequential write.
    with tempfile.TemporaryDirectory() as tmp:
        local = Path(tmp) / path.name
        build_workbook(data).save(local)
        shutil.copyfile(local, path)
    statuses = {}
    for row in data["rows"]:
        statuses.setdefault(row["match_status"], set()).add(row["item_no"])
    print(f"wrote {path}\n  {len(data['rows'])} rows; "
          + ", ".join(f"{k} {len(v)} items" for k, v in sorted(statuses.items())))


if __name__ == "__main__":
    main()
