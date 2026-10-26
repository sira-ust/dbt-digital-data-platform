"""Load the monthly promo workbook into the promo_items landing table (step 1).

TWO LAYOUTS. Since NOV 2026 the month's files (the promo list, and the description
file with photos) are merged into one INPUT workbook in a fixed layout
(scripts/promo/promo_input_schema.py; "UST 2026 NOV PROMO INPUT.xlsx", sheets items + meta).
A workbook with those sheets is read that way: validated first (an error stops the
load), every schema column kept, and each row's photo saved as
<photo dir>/<YYYYMM>/<item_no>.<ext> with its path in photo_path. Any other workbook
is read the original way, below.

The original layout: the workbook arrives by email once a month, e.g. "UST 2026 OCT NEW PROMO.xlsx",
with one sheet per list: "OCT NEW" and "OCT PROMO" (future months follow
"<MON> NEW" / "<MON> PROMO", so sheets are found by SUFFIX, never by name).
The two sheets do not share a layout — OCT NEW has 17 columns including Vendor
No., OCT PROMO has 14 — so columns are read BY HEADER NAME.

Pure from source, like parse_mentions.py: every header is slugified, every value
is written as a string, nothing is cast or cleaned. stg_promo__items does the
typing. Added lineage columns:
    promo_month   first day of the promo month (2026-10-01)
    promo_sheet   the sheet name as written ("OCT NEW")
    sheet_order   1-based position of the sheet in the workbook
    row_order     1-based position of the row within its sheet
    loaded_at / source_file

promo_month = the year in the filename + the month at the start of the sheet
name. If either is missing it falls back to the month of the rows' Ending Date
(the promo runs Sep 27 - Oct 31, so the START date would say September).

Two backends:
  --backend local       writes <local_data_root>/promo/promo_items_<YYYYMM>.parquet
                        (a re-load of the same month replaces that file)
  --backend databricks  appends to ust_databricks.social.promo_items and moves the
                        file to the landing Volume's _archive folder. Staging keeps
                        only the latest load of each month, so a corrected re-drop
                        replaces the earlier one.

Usage:
    python scripts/promo/load_promo_list.py "C:/Users/me/Downloads/UST 2026 OCT NEW PROMO.xlsx"
    python scripts/promo/load_promo_list.py --backend databricks      # every .xlsx in the landing Volume
"""

from __future__ import annotations

import argparse
import glob
import os
import re
import shutil
import sys
from datetime import date, datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path.cwd() / "scripts" / "promo"))
try:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
except NameError:
    pass

from promo_common import SOCIAL_SCHEMA, append_records, ensure_landing, local_data_root  # noqa: E402
from promo_input_schema import COLUMN_NAMES, SHEET_ITEMS, SHEET_META  # noqa: E402

LANDING_DIR = "/Volumes/ust_databricks/social/promo_landing"
ARCHIVE_DIR = f"{LANDING_DIR}/_archive"
PHOTO_DIR_DBX = f"{LANDING_DIR}/_photos"
DBX_TABLE = f"{SOCIAL_SCHEMA}.promo_items"
SHEET_SUFFIXES = ("NEW", "PROMO")
REQUIRED_HEADERS = ["Item No.", "Promo Code", "Description", "Size"]
MONTHS = {m: i for i, m in enumerate(
    ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"], start=1)}
LINEAGE = ["promo_month", "promo_sheet", "sheet_order", "row_order", "loaded_at", "source_file"]


def slugify(header) -> str:
    """'Promo Buy Qty.' -> 'promo_buy_qty'. Same rule as parse_mentions.py."""
    return re.sub(r"[^0-9a-z]+", "_", str(header).strip().lower()).strip("_")


def is_promo_sheet(name: str) -> bool:
    return str(name).strip().upper().split()[-1:] in ([s] for s in SHEET_SUFFIXES)


def month_from_sheet(name: str):
    first = str(name).strip().upper().split()[:1]
    return MONTHS.get(first[0][:3]) if first else None


def year_from_filename(filename: str):
    m = re.search(r"(20\d{2})", filename)
    return int(m.group(1)) if m else None


def _as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.strptime(str(v).strip()[:10], "%Y-%m-%d").date()
    except ValueError:
        try:
            return datetime.strptime(str(v).strip(), "%m/%d/%Y").date()
        except ValueError:
            return None


def resolve_promo_month(filename: str, sheet_names: list[str], ending_dates: list) -> date:
    year = year_from_filename(filename)
    months = {month_from_sheet(s) for s in sheet_names} - {None}
    ends = [d for d in map(_as_date, ending_dates) if d]
    if year and len(months) == 1:
        resolved = date(year, months.pop(), 1)
        if ends:
            latest = max(ends)
            if (latest.year, latest.month) != (resolved.year, resolved.month):
                print(f"  WARNING promo month {resolved:%Y-%m} from the file/sheet names, "
                      f"but rows end {latest:%Y-%m-%d}")
        return resolved
    if ends:
        latest = max(ends)
        return date(latest.year, latest.month, 1)
    raise SystemExit(f"cannot tell the promo month of {filename}: no year in the name, "
                     f"no '<MON> NEW/PROMO' sheet names, no Ending Date")


def is_input_file(path: str) -> bool:
    """The merged INPUT layout (promo_input_schema.py): sheets items AND meta."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return {SHEET_ITEMS, SHEET_META} <= set(wb.sheetnames)
    finally:
        wb.close()


def read_input_file(path: str, photo_dir: Path) -> list[dict]:
    """The rows of a merged INPUT workbook, in the same record shape as
    read_workbook(): promo_sheet is "<MON> <list>" ("NOV NEW"), so the rest of the
    pipeline and the output workbook read it as before. Photos are written to
    photo_dir/<YYYYMM>/<item_no>.<ext>."""
    from validate_promo_input import read_input, validate

    filename = os.path.basename(path)
    data = read_input(path)
    errors, warnings = validate(path, data)
    for w in warnings:
        print(f"  {filename}: warning: {w}")
    if errors:
        raise SystemExit(f"{filename}: not loaded, {len(errors)} schema error(s):\n  "
                         + "\n  ".join(errors))
    month = data["meta"]["promo_month"][0]
    promo_month = date(int(month[:4]), int(month[5:7]), 1)
    mon = [k for k, v in MONTHS.items() if v == promo_month.month][0]
    loaded_at = datetime.now(timezone.utc).replace(tzinfo=None).isoformat(sep=" ", timespec="seconds")
    out_dir = Path(photo_dir) / f"{promo_month:%Y%m}"
    sheet_order, row_order, out = {}, {}, []
    for r in data["rows"]:
        lst = r["promo_list"].upper()
        sheet_order.setdefault(lst, len(sheet_order) + 1)
        row_order[lst] = row_order.get(lst, 0) + 1
        rec = {c: r.get(c) for c in COLUMN_NAMES if c not in ("promo_month", "promo_list", "photo")}
        photo = data["photos"].get(r["_row"])
        rec["photo_path"] = None
        if photo:
            out_dir.mkdir(parents=True, exist_ok=True)
            target = out_dir / f"{r['item_no']}.{photo[1]}"
            target.write_bytes(photo[0])
            rec["photo_path"] = str(target).replace("\\", "/")
        rec.update(promo_month=promo_month.isoformat(), promo_sheet=f"{mon} {lst}",
                   sheet_order=sheet_order[lst], row_order=row_order[lst], loaded_at=loaded_at,
                   source_file=filename)
        out.append(rec)
    print(f"  {filename}: INPUT layout, promo month {promo_month:%Y-%m}, "
          + ", ".join(f"{mon} {k} {v} rows" for k, v in row_order.items())
          + f", {sum(1 for r in out if r['photo_path'])} photos")
    return out


def read_any(path: str, photo_dir: Path) -> list[dict]:
    return read_input_file(path, photo_dir) if is_input_file(path) else read_workbook(path)


def read_workbook(path: str) -> list[dict]:
    """All promo rows of one workbook as string-valued dicts, sheet then row order."""
    import openpyxl

    filename = os.path.basename(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = [ws for ws in wb.worksheets if is_promo_sheet(ws.title)]
    if not sheets:
        raise SystemExit(f"{filename}: no sheet named '<MON> NEW' or '<MON> PROMO' "
                         f"(found {wb.sheetnames})")

    raw = []   # (sheet, sheet_order, row_order, {header: value})
    for sheet_order, ws in enumerate(sheets, start=1):
        rows = list(ws.iter_rows(values_only=True))
        header_idx = next((i for i, r in enumerate(rows[:10]) if r and "Item No." in [
            str(c).strip() if c is not None else None for c in r]), None)
        if header_idx is None:
            raise SystemExit(f"{filename} / {ws.title}: no 'Item No.' header in the first 10 rows")
        headers = [str(c).strip() if c is not None else None for c in rows[header_idx]]
        missing = [h for h in REQUIRED_HEADERS if h not in headers]
        if missing:
            raise SystemExit(f"{filename} / {ws.title}: missing columns {missing}")
        item_col = headers.index("Item No.")
        row_order = 0
        for r in rows[header_idx + 1:]:
            if r is None or item_col >= len(r) or r[item_col] in (None, ""):
                continue   # blank / trailing rows
            row_order += 1
            raw.append((ws.title.strip(), sheet_order, row_order,
                        {h: v for h, v in zip(headers, r) if h}))
    wb.close()

    promo_month = resolve_promo_month(
        filename, [s.title for s in sheets], [vals.get("Ending Date") for *_, vals in raw])
    loaded_at = datetime.now(timezone.utc).replace(tzinfo=None).isoformat(sep=" ", timespec="seconds")

    out = []
    for sheet, sheet_order, row_order, vals in raw:
        rec = {slugify(h): (None if v is None else _stringify(v)) for h, v in vals.items()}
        rec.update(promo_month=promo_month.isoformat(), promo_sheet=sheet, sheet_order=sheet_order,
                   row_order=row_order, loaded_at=loaded_at, source_file=filename)
        out.append(rec)
    by_sheet = {}
    for r in out:
        by_sheet[r["promo_sheet"]] = by_sheet.get(r["promo_sheet"], 0) + 1
    print(f"  {filename}: promo month {promo_month:%Y-%m}, "
          + ", ".join(f"{k} {v} rows" for k, v in by_sheet.items()))
    return out


def _stringify(v) -> str:
    """Excel gives ints as int and dates as datetime; keep them readable as text."""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _columns(records: list[dict]) -> list[str]:
    cols = []
    for r in records:
        for k in r:
            if k not in cols and k not in LINEAGE:
                cols.append(k)
    return cols + LINEAGE


def load_local(paths: list[str], out_root=None) -> None:
    import pandas as pd
    for path in paths:
        recs = read_any(path, local_data_root(out_root) / "promo" / "photos")
        month = recs[0]["promo_month"][:7].replace("-", "")
        out = local_data_root(out_root) / "promo" / f"promo_items_{month}.parquet"
        out.parent.mkdir(parents=True, exist_ok=True)
        pd.DataFrame(recs, columns=_columns(recs)).astype(object).where(
            lambda d: d.notna(), None).to_parquet(out, index=False)
        print(f"  wrote {len(recs)} rows -> {out}")


def load_databricks() -> None:
    files = sorted(glob.glob(os.path.join(LANDING_DIR, "*.xlsx")))
    print(f"{len(files)} xlsx file(s) in {LANDING_DIR}")
    for path in files:
        recs = read_any(path, Path(PHOTO_DIR_DBX))
        append_records("databricks", recs, None, DBX_TABLE, _columns(recs))
        os.makedirs(ARCHIVE_DIR, exist_ok=True)
        shutil.move(path, os.path.join(ARCHIVE_DIR, os.path.basename(path)))
        print(f"  appended {len(recs)} rows to {DBX_TABLE}; archived {os.path.basename(path)}")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("paths", nargs="*", help="promo workbook(s) (local backend)")
    ap.add_argument("--backend", choices=["local", "databricks"], default="local")
    ap.add_argument("--out-root", help="override local_data_root (e.g. data/real)")
    args = ap.parse_args(argv)
    if args.backend == "databricks":
        load_databricks()
    else:
        if not args.paths:
            ap.error("give the promo workbook path(s)")
        load_local(args.paths, args.out_root)
    # first step of the monthly run, so the downstream landing tables exist
    # before dbt builds the staging views over them
    ensure_landing(args.backend, args.out_root)


if __name__ == "__main__":
    main()
