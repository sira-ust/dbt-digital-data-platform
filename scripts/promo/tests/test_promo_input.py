"""The monthly promo INPUT file: schema, validator and loader, on small synthetic
workbooks. Run:  python -m pytest scripts/promo/tests
"""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

import load_promo_list as loader  # noqa: E402
import validate_promo_input as v  # noqa: E402
from promo_input_schema import COLUMN_NAMES, SCHEMA_VERSION  # noqa: E402

PROMPT = Path(__file__).resolve().parents[1] / "prompts" / "promo_merge_input.md"


def _row(item_no, desc, **kw):
    base = {c: None for c in COLUMN_NAMES}
    base.update(promo_month="2026-11", promo_list="NEW", item_no=item_no, description=desc, size="12X6OZ")
    base.update(kw)
    return base


def _png():
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "red").save(buf, format="PNG")
    buf.seek(0)
    return buf


def make(tmp_path, rows, name="UST 2026 NOV PROMO INPUT.xlsx", meta=None, header=None, photo_rows=(),
         unmatched=0):
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "items"
    header = header or COLUMN_NAMES
    ws.append(header)
    for r in rows:
        ws.append([r.get(c) for c in header])
    for n in photo_rows:                                   # a picture anchored on the photo cell
        ws.add_image(XLImage(_png()), f"{get_column_letter(header.index('photo') + 1)}{n}")
    m = wb.create_sheet("meta")
    m.append(["key", "value"])
    for k, val in (meta or [("schema_version", SCHEMA_VERSION), ("promo_month", "2026-11"),
                            ("merged_on", "2026-10-27"), ("source", "UST 2026 NOV NEW PROMO.xlsx (promo_list, 2 rows)")]):
        m.append([k, val])
    if unmatched:
        u = wb.create_sheet("unmatched_detail")
        u.append(["No.", "Description"])
        for i in range(unmatched):
            u.append([90000 + i, "SOMETHING"])
    path = tmp_path / name
    wb.save(path)
    return str(path)


GOOD = [_row("35060", "HOSHI GUMMY MANGO", origin_code="TH", eta="2026-10-17",
             selling_points="Peelable Fruit Skin | Two Different Textures"),
        _row("35061", "HOSHI GUMMY MANGO", size="24X3.53OZ")]


def test_valid_file(tmp_path):
    errors, warnings = v.validate(make(tmp_path, GOOD, photo_rows=[2]))
    assert errors == []
    assert any("2 rows (NEW 2); 1 with a description or selling points, 1 with a photo" in w for w in warnings)


def test_prompt_names_every_column():
    """The merge prompt and the schema must not drift apart."""
    text = PROMPT.read_text(encoding="utf-8")
    assert all(f"`{c}`" in text for c in COLUMN_NAMES)
    assert f"`schema_version` | `{SCHEMA_VERSION}`" in text


@pytest.mark.parametrize("change, message", [
    ({"item_no": "35060.0"}, "is not digits only"),
    ({"item_no": None}, "item_no is blank"),
    ({"size": None}, "size is blank"),
    ({"promo_list": "SALE"}, "is not one of"),
    ({"eta": "10/17/2026"}, "is not a date"),
    ({"promo_month": "2026-12"}, "differs from meta"),
])
def test_row_errors(tmp_path, change, message):
    rows = [{**GOOD[0], **change}, GOOD[1]]
    errors, _ = v.validate(make(tmp_path, rows))
    assert any(message in e for e in errors), errors


def test_same_item_twice_on_a_list_is_an_error_but_on_both_lists_is_not(tmp_path):
    errors, _ = v.validate(make(tmp_path, [GOOD[0], GOOD[0]]))
    assert any("NEW item 35060 is already on row 2" in e for e in errors)
    errors, _ = v.validate(make(tmp_path, [GOOD[0], {**GOOD[0], "promo_list": "PROMO"}]))
    assert errors == []


def test_missing_required_column(tmp_path):
    header = [c for c in COLUMN_NAMES if c != "size"]
    errors, _ = v.validate(make(tmp_path, GOOD, header=header))
    assert any("missing required columns ['size']" in e for e in errors)


def test_optional_and_unknown_columns_only_warn(tmp_path):
    header = ["promo_month", "promo_list", "item_no", "description", "size", "colour"]
    errors, warnings = v.validate(make(tmp_path, GOOD, header=header))
    assert errors == []
    assert any("not in the schema are ignored: ['colour']" in w for w in warnings)
    assert any("optional columns not present" in w for w in warnings)


def test_meta_and_file_name(tmp_path):
    bad = [("schema_version", "2"), ("promo_month", "2026-11"), ("merged_on", "2026-10-27"), ("source", "x")]
    errors, _ = v.validate(make(tmp_path, GOOD, meta=bad))
    assert any("schema_version 2" in e for e in errors)
    errors, _ = v.validate(make(tmp_path, GOOD, name="UST 2026 DEC PROMO INPUT.xlsx"))
    assert any("file name says 2026 DEC" in e for e in errors)
    errors, _ = v.validate(make(tmp_path, GOOD, meta=[("schema_version", "1"), ("promo_month", "2026-11")]))
    assert any("no 'merged_on' row" in e for e in errors) and any("no 'source' row" in e for e in errors)


def test_selling_points_must_be_pipe_joined(tmp_path):
    rows = [{**GOOD[0], "selling_points": "Peelable Fruit Skin\nTwo Different Textures"}, GOOD[1]]
    errors, warnings = v.validate(make(tmp_path, rows))
    assert errors == [] and any("joined with ' | '" in w for w in warnings)


def test_unmatched_detail_is_reported(tmp_path):
    _, warnings = v.validate(make(tmp_path, GOOD, unmatched=3))
    assert any("3 description-file row(s) are on no promo list" in w for w in warnings)


def test_loader_reads_the_input_layout(tmp_path):
    path = make(tmp_path, GOOD, photo_rows=[3])
    assert loader.is_input_file(path)
    recs = loader.read_input_file(path, tmp_path / "photos")
    assert [(r["promo_sheet"], r["row_order"], r["item_no"]) for r in recs] == [
        ("NOV NEW", 1, "35060"), ("NOV NEW", 2, "35061")]
    assert recs[0]["promo_month"] == "2026-11-01" and recs[0]["size"] == "12X6OZ"
    assert recs[0]["selling_points"] == "Peelable Fruit Skin | Two Different Textures"
    # the photo on the second row is saved under the month, by item number
    assert recs[0]["photo_path"] is None
    saved = Path(recs[1]["photo_path"])
    assert saved.name == "35061.png" and saved.parent.name == "202611" and saved.read_bytes()[:4] == b"\x89PNG"
    assert "photo" not in recs[0] and "promo_list" not in recs[0]


def test_loader_refuses_an_invalid_file(tmp_path):
    path = make(tmp_path, [{**GOOD[0], "item_no": "abc"}])
    with pytest.raises(SystemExit, match="not loaded, 1 schema error"):
        loader.read_input_file(path, tmp_path / "photos")


def test_an_original_promo_workbook_is_still_read_the_old_way(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "NOV NEW"
    ws.append(["Item No.", "Promo Code", "Description", "Size"])
    ws.append([35060, None, "HOSHI GUMMY MANGO", "12X17.64OZ"])
    path = tmp_path / "UST 2026 NOV NEW PROMO.xlsx"
    wb.save(path)
    assert not loader.is_input_file(str(path))
    recs = loader.read_any(str(path), tmp_path / "photos")
    assert recs[0]["promo_sheet"] == "NOV NEW" and recs[0]["item_no"] == "35060"


def test_floating_photos_are_placed_in_their_cells(tmp_path):
    """photos_in_cells.py: floating photos become in-cell pictures on the same rows, with
    the same bytes; the validator warns about floating ones; a second run changes nothing."""
    import photos_in_cells
    path = make(tmp_path, GOOD, photo_rows=[2, 3])
    before = v.read_input(path)
    _, warnings = v.validate(path, before)
    assert before["floating_photos"] == 2 and any("float over their cells" in w for w in warnings)

    assert photos_in_cells.place_photos_in_cells(path) == 2
    after = v.read_input(path)
    assert after["floating_photos"] == 0
    assert sorted(after["photos"]) == [2, 3]
    assert {r: d for r, (d, _) in after["photos"].items()} == {r: d for r, (d, _) in before["photos"].items()}
    assert v._in_cell_pictures(path, "items", "Q").keys() == {2, 3}      # read as in-cell pictures
    errors, warnings = v.validate(path, after)
    assert errors == [] and not any("float over" in w for w in warnings)
    assert [r["item_no"] for r in after["rows"]] == ["35060", "35061"]    # the data is untouched
    assert photos_in_cells.place_photos_in_cells(path) == 0                # nothing left to move


def _result(tmp_path, header, items):
    """A finished output workbook's main sheet, as far as add_photos_to_output needs it."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Promo vs Social Listening"
    ws.append(header)
    for no in items:
        ws.append([no if h == "Item No." else ("=1+1" if h == "Posts" else None) for h in header])
    path = tmp_path / "result.xlsx"
    wb.save(path)
    return str(path)


def test_photos_are_added_to_an_older_output_by_item_number(tmp_path):
    """An output exported before the Photo column: the column is added at the end,
    each photo lands on its own item's row (35061 and 35060 share a description),
    formulas stay, and a second run adds nothing."""
    import add_photos_to_output as ap
    src = make(tmp_path, GOOD, photo_rows=[2, 3])          # 35060 on row 2, 35061 on row 3
    expected = {r["item_no"]: v.read_input(src)["photos"][r["_row"]][0] for r in v.read_input(src)["rows"]}
    result = _result(tmp_path, ["Promo Sheet", "Item No.", "Promo Item", "Posts", "Trending Social 1"],
                     [35061, 99999, 35060])
    assert ap.add_photos(result, src) == (2, "F")
    ws = load_workbook(result)["Promo vs Social Listening"]
    assert ws["F1"].value == "Photo" and ws["D2"].value == "=1+1"
    by_row = {i.anchor._from.row + 1: i._data() for i in ws._images}
    assert by_row == {2: expected["35061"], 4: expected["35060"]}   # 99999 has no photo
    assert ap.add_photos(result, src) == (0, "F")


def test_the_filter_takes_in_the_photo_column_and_rows_match(tmp_path):
    """A Photo column outside the header filter is left behind by a sort from the
    filter arrows (NOV 2026: the photos stayed put while the items moved). Both tools
    widen the filter to it, and give every item row one height so a sorted photo never
    lands on a short row. The in-cell step keeps both."""
    import add_photos_to_output as ap
    import photos_in_cells
    src = make(tmp_path, GOOD, photo_rows=[2, 3])
    wb = Workbook()
    ws = wb.active
    ws.title = "Promo vs Social Listening"
    ws.append(["Promo Sheet", "Item No.", "Promo Item", "Posts", "Trending Social 1"])
    for no in (35061, 99999, 35060):
        ws.append(["NOV NEW", no, "x", "=1+1", None])
    ws.auto_filter.ref = "A1:E4"
    result = str(tmp_path / "result.xlsx")
    wb.save(result)

    ap.add_photos(result, src)
    ws = load_workbook(result)["Promo vs Social Listening"]
    assert ws.auto_filter.ref == "A1:F4"
    assert {ws.row_dimensions[r].height for r in (2, 3, 4)} == {ap.PHOTO_ROW_HEIGHT}   # 99999 too

    assert photos_in_cells.place_photos_in_cells(result, sheet="Promo vs Social Listening", column="Photo") == 2
    assert v._in_cell_pictures(result, "Promo vs Social Listening", "F").keys() == {2, 4}
    import zipfile
    with zipfile.ZipFile(result) as z:
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        book = z.read("xl/workbook.xml").decode()
    assert '<autoFilter ref="A1:F4"' in sheet and "$A$1:$F$4" in book        # the hidden filter name too
    ws = load_workbook(result)["Promo vs Social Listening"]
    assert {ws.row_dimensions[r].height for r in (2, 3, 4)} == {ap.PHOTO_ROW_HEIGHT}


def test_photos_go_in_an_existing_photo_column(tmp_path):
    import add_photos_to_output as ap
    src = make(tmp_path, GOOD, photo_rows=[2])
    result = _result(tmp_path, ["Promo Sheet", "Item No.", "Promo Item", "Photo", "Size"], [35060])
    assert ap.add_photos(result, src) == (1, "D")


def test_in_cell_pictures_are_read():
    """Real 'Place in Cell' pictures (Excel rich data, invisible to openpyxl), from the
    NOV description file, when it is on this machine: 29 photos in its Image column (G)."""
    src = Path("C:/Users/sirasasitorn/Downloads/UST 2026 NOV NEW - desc.xlsx")
    if not src.exists():
        pytest.skip("NOV description file not on this machine")
    title = load_workbook(src, read_only=True).sheetnames[0]
    photos = v._in_cell_pictures(str(src), title, "G")
    assert len(photos) == 29 and all(ext in ("png", "jpeg") for _, ext in photos.values())
