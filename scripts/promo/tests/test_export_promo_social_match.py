"""The workbook export: layout, formatting and live formulas, from synthetic rows.

No database needed — build_workbook() takes the same dict read_inputs() returns.
The formula tests need the optional `formulas` package (a spreadsheet engine) and
are skipped without it. Run:  python -m pytest scripts/promo/tests
"""

import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import export_promo_social_match as ex  # noqa: E402

MAIN = "Promo vs Social Listening"


def _row(sheet, row_order, item_no, desc, status, rank=1, **match):
    base = {"promo_sheet": sheet, "row_order": row_order, "match_rank": rank, "item_no": item_no,
            "item_no_int": int(item_no), "promo_description": desc + "  ", "promo_size": "24X10.9OZ",
            "promo_code": None if sheet.endswith("NEW") else "5+1", "match_status": status,
            "matched_content_item": None, "channel": None, "profile": None, "post_date": None,
            "engagement_text": None, "caption": None, "link": None, "source": None,
            "search_language": None, "match_note": None, "views": None, "likes": None}
    base.update(match)
    # the mart's item-level columns; one source per item unless a test overrides them
    src = base["source"]
    base.setdefault("item_match_status", status)
    base.setdefault("social_match_status", status if src == "Social Listening" else None)
    base.setdefault("scraped_match_status", status if src == "Scraped" else None)
    return base


def _item(sheet, row_order, item_no, desc, branded=True, **kw):
    base = {"promo_sheet": sheet, "row_order": row_order, "item_no": item_no, "promo_description": desc,
            "promo_size": "24X10.9OZ", "is_branded": branded, "brand": "Fanta" if branded else None,
            "product": "soft drink", "flavor_variant": "grape", "search_queries": "แฟนต้า องุ่น|fanta grape",
            "brand_terms": "fanta|แฟนต้า", "flavor_terms": "grape|องุ่น", "negative_terms": None,
            "source_file": "UST 2026 OCT NEW PROMO.xlsx"}
    base.update(kw)
    return base


SL = dict(matched_content_item="ดีโด้ น้ำส้มสายน้ำผึ้ง (Deedo Sainamphueng orange juice)",
          channel="Web", profile="trueid.net", post_date="2026-07-13", engagement_text="Not reported",
          caption="มิกซ์ กาแฟส้ม", link="https://www.trueid.net/watch/th-th/short/8yXZnXjx325b",
          source="Social Listening", match_note="Brand, product and flavor match.", views=0, likes=0)
YT = dict(matched_content_item="Fanta Grape", channel="YouTube", profile="Fanta TH", post_date="2022-03-14",
          engagement_text="176,000 views, 900 likes", caption="ระเบิดความซ่ากับองุ่นป๊อบ!",
          link="https://www.youtube.com/watch?v=GUwN8yYJboI", source="Scraped", search_language="Thai",
          match_note="Older than 2 years: no recent video was found. Brand + product/flavor named in video.",
          views=176000, likes=900)


@pytest.fixture
def data():
    rows = [
        _row("OCT NEW", 1, "86251", "BAMBOO SUSHI MAT", "Unmatched"),
        _row("OCT NEW", 2, "37353", "BENTO SQUID SNACK (BBQ CHEESE)", "Close",
             **{**SL, "matched_content_item": "Mực Bento", "match_note": "Same brand/product line"}),
        _row("OCT PROMO", 1, "34585", "DEEDO SAINAMPHUENG FRUIT JUICE", "Exact", **SL),
        # Fanta Grape: one video from the Thai phrase, one from the English phrase
        _row("OCT PROMO", 2, "34662", "FANTA GRAPE FLAVOR", "Exact", 1, **YT),
        _row("OCT PROMO", 2, "34662", "FANTA GRAPE FLAVOR", "Exact", 2,
             **{**YT, "link": "https://youtu.be/x", "search_language": "English",
                "engagement_text": "505 views", "post_date": "2025-06-03",
                "match_note": "Brand + product/flavor named in video."}),
        _row("OCT PROMO", 3, "34658", "FANTA ORANGE FLAVOR", "Unmatched"),
    ]
    items = [_item("OCT NEW", 1, "86251", "BAMBOO SUSHI MAT", branded=False),
             _item("OCT NEW", 2, "37353", "BENTO SQUID SNACK (BBQ CHEESE)"),
             _item("OCT PROMO", 1, "34585", "DEEDO SAINAMPHUENG FRUIT JUICE"),
             _item("OCT PROMO", 2, "34662", "FANTA GRAPE FLAVOR"),
             _item("OCT PROMO", 3, "34658", "FANTA ORANGE FLAVOR", flavor_terms=None)]
    return {"promo_month": "2026-10-01", "rows": rows, "items": items,
            "mentions": {"n": 53601, "first_at": datetime(2026, 2, 20), "last_at": datetime(2026, 9, 29)},
            "searched_at": datetime(2026, 10, 2, 5, 0)}


@pytest.fixture
def saved(data, tmp_path):
    path = tmp_path / ex.output_name(data["promo_month"])
    ex.build_workbook(data).save(path)
    return path


def _calc(path):
    """Recalculate a saved workbook; -> a lookup of 'Sheet'!A1 -> value."""
    formulas = pytest.importorskip("formulas")
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()

    def val(sheet, ref):
        key = next(k for k in sol if k.upper().endswith(f"{sheet.upper()}'!{ref}"))
        v = sol[key].value
        return float(v[0][0] if hasattr(v, "__getitem__") else v)
    return val


def test_file_name(data):
    assert ex.output_name(data["promo_month"]) == "UST_OCT_Promo_Social_Listening_Match.xlsx"


def test_sheets_and_header(saved):
    wb = load_workbook(saved)
    assert wb.sheetnames == ["Summary", MAIN, "TikTok Search List", "RedNote Search List"]
    assert wb.active.title == "Summary"                                   # opens on the Summary
    ws = wb[MAIN]
    # the item, then a block of 3 per source: Trending Social, then each language searched
    # (English and Thai here, from the Fanta items' phrases), in a fixed order
    assert [c.value for c in ws[1]] == [
        "Promo Sheet", "Item No.", "OCT Promo Item", "Size", "Promo Code", "Match Status",
        "On Trending Social (Mentionlytics)", "Direct Scraping (TikTok, YouTube & RedNote)", "Posts",
        "Trending Social 1", "Trending Social 2", "Trending Social 3",
        "English 1", "English 2", "English 3", "Thai 1", "Thai 2", "Thai 3"]
    h = ws["A1"]
    assert h.font.bold and h.font.name == "Arial" and h.font.color.rgb == "FFFFFFFF"
    assert h.fill.fgColor.rgb == "FF1F4E78" and h.alignment.vertical == "center" and h.alignment.wrap_text
    assert ws.row_dimensions[1].height == 27.75
    # blocks' headers alternate, so they read apart
    assert [ws[f"{c}1"].fill.fgColor.rgb for c in "JMP"] == ["FF1F4E78", "FF2E75B6", "FF1F4E78"]


def test_layout(saved):
    ws = load_workbook(saved)[MAIN]
    widths = [ws.column_dimensions[c].width for c in "ABCDEFGHIJKLMNOPQR"]
    assert widths == [13, 10, 32, 15, 10, 12, 16, 16, 8] + [27] * 9
    assert ws.freeze_panes == "D2"
    assert ws.auto_filter.ref == "A1:R6"
    assert all(ws.row_dimensions[r].height == 39.75 for r in range(2, ws.max_row + 1))


def test_one_row_per_item(saved):
    ws = load_workbook(saved)[MAIN]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [r[1] for r in rows] == [86251, 37353, 34585, 34662, 34658]      # Fanta Grape once
    assert [r[0] for r in rows] == ["OCT NEW"] * 2 + ["OCT PROMO"] * 3
    assert rows[0][2] == "BAMBOO SUSHI MAT"                                 # trimmed
    assert rows[0][4] is None                                               # blank promo code
    body = ws["C3"]
    assert body.font.name == "Arial" and body.font.size == 10
    assert body.alignment.vertical == "top" and body.alignment.wrap_text
    assert body.border.left.color.rgb == "FFBFBFBF"
    assert ws["C2"].font.color.rgb == "FF808080"                            # Unmatched item greyed


def test_each_language_has_its_own_block(saved):
    ws = load_workbook(saved)[MAIN]
    # Fanta Grape (row 5): the English video and the Thai one, each in its own block
    assert ws["M5"].value == "YouTube · 505 views · 2025-06-03"
    assert ws["M5"].hyperlink.target == "https://youtu.be/x"
    assert ws["P5"].value == "YouTube · 176,000 views · 900 likes · 2022-03-14 · older"
    assert ws["P5"].hyperlink.target.startswith("https://www.youtube.com/watch")
    assert ws["J5"].value is None                                           # no social post


def test_result_cells_style_and_note(saved):
    ws = load_workbook(saved)[MAIN]
    cell = ws["P5"]                                                          # Exact: green
    assert cell.fill.fgColor.rgb == "FFC6EFCE"
    assert cell.font.underline == "single" and cell.font.color.rgb == "FF0563C1"
    assert "ระเบิดความซ่ากับองุ่นป๊อบ!" in cell.comment.text and "Older than 2 years" in cell.comment.text
    bento = ws["J3"]                                                         # Close: yellow
    assert bento.value == "Web · Not reported · 2026-07-13"
    assert bento.fill.fgColor.rgb == "FFFFEB9C"
    assert bento.hyperlink.target.startswith("https://www.trueid.net/")


def test_not_searched_vs_nothing_found(saved):
    ws = load_workbook(saved)[MAIN]
    # sushi mat: unbranded, never searched; Bento / Deedo: social listening matched,
    # so not searched: "–" opens each language block
    for r in (2, 3, 4):
        assert (ws[f"M{r}"].value, ws[f"P{r}"].value) == ("–", "–")
        assert ws[f"N{r}"].value is None
    # Fanta Orange: searched in English and Thai, nothing found: blank
    assert (ws["M6"].value, ws["P6"].value) == (None, None)
    # the social block is never "not searched": every item is matched against it
    assert ws["J2"].value is None


def test_status_columns_and_banding(saved):
    ws = load_workbook(saved)[MAIN]
    styles = {ws[f"F{r}"].value: (ws[f"F{r}"].fill.fgColor.rgb, ws[f"F{r}"].font.color.rgb)
              for r in range(2, ws.max_row + 1)}
    assert styles == {"Unmatched": ("FFEDEDED", "FF595959"), "Close": ("FFFFEB9C", "FF9C5700"),
                      "Exact": ("FFC6EFCE", "FF006100")}
    assert (ws["G4"].value, ws["H4"].value) == ("Exact", None)              # Deedo: social only
    assert (ws["G5"].value, ws["H5"].value) == (None, "Exact")              # Fanta Grape: scraped only
    shaded = [ws[f"C{r}"].fill.fgColor.rgb == "FFEEF3F8" for r in range(2, ws.max_row + 1)]
    assert shaded == [False, True, False, True, False]                      # every other item


def test_posts_is_a_live_count(saved):
    ws = load_workbook(saved)[MAIN]
    assert ws["I5"].value == '=COUNTA(J5:R5)-COUNTIF(J5:R5,"–")'


def test_more_than_three_results_in_a_language_are_capped_but_another_language_keeps_its_own(data, tmp_path):
    """The point of the blocks: 3 English videos do not push the Thai one out."""
    en = [_row("OCT PROMO", 2, "34662", "FANTA GRAPE FLAVOR", "Exact", n,
               **{**YT, "search_language": "English", "link": f"https://youtu.be/{n}",
                  "engagement_text": f"{n},000 views"}) for n in range(1, 4)]
    th = _row("OCT PROMO", 2, "34662", "FANTA GRAPE FLAVOR", "Exact", 4,
              **{**YT, "search_language": "Thai", "link": "https://youtu.be/th", "engagement_text": "12 views"})
    social = [_row("OCT PROMO", 1, "34585", "DEEDO SAINAMPHUENG FRUIT JUICE", "Exact", n,
                   **{**SL, "link": f"https://t.co/{n}"}) for n in range(1, 6)]
    data["rows"] = social + en + [th]
    data["items"] = [_item("OCT PROMO", 1, "34585", "DEEDO SAINAMPHUENG FRUIT JUICE"),
                     _item("OCT PROMO", 2, "34662", "FANTA GRAPE FLAVOR",
                           search_queries="fanta grape|แฟนต้า องุ่น")]
    path = tmp_path / "blocks.xlsx"
    ex.build_workbook(data).save(path)
    ws = load_workbook(path)[MAIN]
    header = [c.value for c in ws[1]]
    assert header[9:] == ["Trending Social 1", "Trending Social 2", "Trending Social 3",
                          "English 1", "English 2", "English 3", "Thai 1", "Thai 2", "Thai 3"]
    assert [ws.cell(3, c).value for c in (13, 14, 15)] == [
        "YouTube · 1,000 views · 2022-03-14 · older", "YouTube · 2,000 views · 2022-03-14 · older",
        "YouTube · 3,000 views · 2022-03-14 · older"]
    assert ws.cell(3, 16).value == "YouTube · 12 views · 2022-03-14 · older"
    # five social posts: three cells, the third says how many more there are
    assert ws.cell(2, 12).value.endswith("(+2 more)")


def test_summary_formulas_and_footnote(saved):
    sm = load_workbook(saved)["Summary"]
    assert [c.value for c in sm[1]][:3] == ["Match Status", "Promo Items", "Matched Posts"]
    assert [sm[f"A{r}"].value for r in range(2, 6)] == ["Exact", "Close", "Unmatched", "Total"]
    assert sm["B2"].value == f"=COUNTIF('{MAIN}'!$F$2:$F$5000,A2)"
    assert sm["C2"].value == f"=SUMIF('{MAIN}'!$F$2:$F$5000,A2,'{MAIN}'!$I$2:$I$5000)"
    assert sm["B5"].value == "=SUM(B2:B4)" and sm["C5"].value == "=SUM(C2:C4)"
    note = sm["A7"].value
    assert "UST 2026 OCT NEW PROMO.xlsx (5 items)" in note and "53,601 posts" in note
    assert "Oct 2, 2026" in note and "– = not searched in that language" in note


def test_summary_sections(saved):
    sm = load_workbook(saved)["Summary"]
    col_a = [sm.cell(r, 1).value for r in range(1, sm.max_row + 1)]
    assert sm["A9"].value == "Promo items by sheet"
    assert [c.value for c in sm[10]][:4] == ["Match Status", "OCT NEW", "OCT PROMO", "Total"]
    assert sm["B11"].value == f"=COUNTIFS('{MAIN}'!$A$2:$A$5000,\"OCT NEW\",'{MAIN}'!$F$2:$F$5000,$A11)"
    assert sm["D11"].value == "=SUM(B11:C11)" and sm["A14"].value == "Total"
    assert sm["A16"].value == "Matched posts by channel"
    # channels in the results, plus TikTok and RedNote (their hand steps add them), alphabetical
    assert col_a[17:22] == ["RedNote", "TikTok", "Web", "YouTube", "Total"]
    assert sm["B18"].value == f"=COUNTIF('{MAIN}'!$J$2:$R$5000,$A18&\" ·*\")"
    i = col_a.index("Matched posts by search language") + 1
    assert [sm.cell(i + n, 1).value for n in range(2, 6)] == ["Trending Social", "English", "Thai", "Total"]
    i = col_a.index("Unmatched promo items (no post found)") + 1
    assert [c.value for c in sm[i + 1]][:3] == ["Item No.", "OCT Promo Item", "Promo Sheet"]
    assert [sm.cell(i + 2, 1).value, sm.cell(i + 3, 1).value] == [86251, 34658]
    assert sm.cell(i + 4, 1).value.startswith("List captured when the workbook was built")


def test_summary_formulas_recalculate(saved):
    val = _calc(saved)
    # Exact: Deedo (1 post) + Fanta Grape (2) = 2 items, 3 posts; Close: Bento 1/1;
    # Unmatched: 2 items, 0 posts ("–" cells are not posts)
    assert [val("Summary", f"B{r}") for r in (2, 3, 4, 5)] == [2, 1, 2, 5]
    assert [val("Summary", f"C{r}") for r in (2, 3, 4, 5)] == [3, 1, 0, 4]
    # by sheet: OCT NEW has the sushi mat (Unmatched) and Bento (Close); OCT PROMO the rest
    assert [val("Summary", f"B{r}") for r in (11, 12, 13)] == [0, 1, 1]
    assert [val("Summary", f"C{r}") for r in (11, 12, 13)] == [2, 0, 1]
    # by channel: RedNote 0, TikTok 0, Web 2 (Bento, Deedo), YouTube 2;
    # by block: social 2, English 1, Thai 1
    assert [val("Summary", f"B{r}") for r in (18, 19, 20, 21, 22)] == [0, 0, 2, 2, 4]
    assert [val("Summary", f"B{r}") for r in (26, 27, 28, 29)] == [2, 1, 1, 4]


def test_a_result_added_by_hand_is_counted(saved):
    """What the TikTok step does: write a result into an empty cell. The counts follow."""
    wb = load_workbook(saved)
    wb[MAIN]["M6"] = "TikTok · 3,706 likes · 2025-02-14"                  # Fanta Orange, English 1
    wb.save(saved)
    val = _calc(saved)
    assert val(MAIN, "I6") == 1
    assert val("Summary", "B19") == 1                                        # TikTok
    assert val("Summary", "B27") == 2                                        # English


def test_tiktok_list_has_unmatched_branded_items_only(saved):
    tk = load_workbook(saved)["TikTok Search List"]
    items = [r[0] for r in tk.iter_rows(min_row=2, values_only=True)]
    # unbranded sushi mat: no; Bento + Deedo matched in social listening: no;
    # Fanta Grape (YouTube Exact) and Fanta Orange (Unmatched): yes
    assert items == [34662, 34658]
    assert tk["G2"].value == "Exact (YouTube)" and tk["G3"].value == "Unmatched"
    assert tk["H2"].value == "[Thai] แฟนต้า องุ่น\n[English] fanta grape"     # tagged with its block
    assert tk["J3"].value == "(no flavor to require)"
    assert tk["L2"].hyperlink.target.startswith("https://www.tiktok.com/search/video?q=")


def test_chinese_is_searched_only_on_rednote(data, tmp_path):
    """Chinese phrases go to the RedNote list and nowhere else; the other languages to
    TikTok (and YouTube) only. An item with no Chinese phrase is not on the RedNote list."""
    data["items"] = [_item("OCT PROMO", 1, "34662", "FANTA GRAPE FLAVOR",
                           search_queries="แฟนต้า องุ่น|fanta grape|芬达 葡萄"),
                     _item("OCT PROMO", 2, "34658", "FANTA ORANGE FLAVOR")]
    data["rows"] = [_row("OCT PROMO", 1, "34662", "FANTA GRAPE FLAVOR", "Unmatched"),
                    _row("OCT PROMO", 2, "34658", "FANTA ORANGE FLAVOR", "Unmatched")]
    path = tmp_path / "rednote.xlsx"
    ex.build_workbook(data).save(path)
    wb = load_workbook(path)
    tk, rn = wb["TikTok Search List"], wb["RedNote Search List"]
    assert [r[0] for r in tk.iter_rows(min_row=2, values_only=True)] == [34662, 34658]
    assert tk["H2"].value == "[Thai] แฟนต้า องุ่น\n[English] fanta grape"
    assert [r[0] for r in rn.iter_rows(min_row=2, values_only=True)] == [34662]
    assert rn["L1"].value == "RedNote Search Link"
    assert rn["H2"].value == "[Chinese] 芬达 葡萄"
    assert rn["L2"].hyperlink.target.startswith("https://www.xiaohongshu.com/search_result?keyword=%E8%8A%AC")
    # the Chinese block is still there, blank, for the RedNote step to fill
    header = [c.value for c in wb[MAIN][1]]
    assert wb[MAIN].cell(2, header.index("Chinese 1") + 1).value is None
    assert ex.platform_phrases("fanta grape|芬达 葡萄|ハタ ソーダ", "youtube") == ["fanta grape"]
    assert ex.platform_phrases("fanta grape|芬达 葡萄|ハタ ソーダ", "rednote") == ["芬达 葡萄"]


def test_own_brand_shows_both_sources(data, tmp_path):
    """An own-brand item keeps its social listening post AND gets scraped: one item
    status (the better one), each source's result in its own block and column, and it
    stays on the TikTok list although social listening matched it."""
    sl = dict(matched_content_item="Bánh tráng Chuồn Chuồn", channel="TikTok", profile="@a",
              post_date="2026-09-01", engagement_text="12 likes", caption="chuồn chuồn",
              link="https://www.tiktok.com/@a/video/1", source="Social Listening",
              match_note="Same brand/product line", views=0, likes=12)
    yt = dict(matched_content_item="Chuồn Chuồn rice paper", channel="YouTube", profile="Chuon Chuon Eats",
              post_date="2026-06-03", engagement_text="505 views", caption="bánh tráng",
              link="https://www.youtube.com/watch?v=y", source="Scraped", search_language="Vietnamese",
              match_note="Brand + product/flavor named in video.", views=505, likes=5)
    both = dict(item_match_status="Exact", social_match_status="Close", scraped_match_status="Exact")
    data["rows"] = [_row("OCT NEW", 1, "43220", "CHUON CHUON RICE PAPER", "Close", 1, **sl, **both),
                    _row("OCT NEW", 1, "43220", "CHUON CHUON RICE PAPER", "Exact", 2, **yt, **both)]
    data["items"] = [_item("OCT NEW", 1, "43220", "CHUON CHUON RICE PAPER", brand="Chuon Chuon",
                           is_own_brand=True, search_queries="bánh tráng chuồn chuồn|dragonfly rice paper")]
    path = tmp_path / "own.xlsx"
    ex.build_workbook(data).save(path)
    wb = load_workbook(path)
    ws = wb[MAIN]
    assert ws.max_row == 2                                                   # one row for the item
    assert (ws["F2"].value, ws["G2"].value, ws["H2"].value) == ("Exact", "Close", "Exact")
    header = [c.value for c in ws[1]]
    assert ws.cell(2, header.index("Trending Social 1") + 1).value == "TikTok · 12 likes · 2026-09-01"
    assert ws.cell(2, header.index("Vietnamese 1") + 1).value == "YouTube · 505 views · 2026-06-03"
    assert ws.cell(2, header.index("English 1") + 1).value is None            # searched, nothing found
    tk = wb["TikTok Search List"]
    assert tk["A2"].value == 43220
    assert tk["G2"].value == "Close (On Trending Social); Exact (YouTube)"


def test_frozen_item_with_no_brand_is_on_the_tiktok_list(data, tmp_path):
    """A frozen item with no brand (FZ) is searched as a frozen product: it is on the
    TikTok list, with the words for frozen standing in for the brand."""
    data["items"] = [_item("OCT NEW", 1, "73094", "FZ PURE SUGARCANE JUICE", branded=False,
                           is_frozen=True, flavor_terms=None,
                           search_queries="nước mía đông lạnh|frozen sugarcane juice|冷冻甘蔗汁"),
                     _item("OCT NEW", 2, "86251", "BAMBOO SUSHI MAT", branded=False, is_frozen=False)]
    data["rows"] = [_row("OCT NEW", 1, "73094", "FZ PURE SUGARCANE JUICE", "Unmatched"),
                    _row("OCT NEW", 2, "86251", "BAMBOO SUSHI MAT", "Unmatched")]
    path = tmp_path / "frozen.xlsx"
    ex.build_workbook(data).save(path)
    wb = load_workbook(path)
    tk = wb["TikTok Search List"]
    rows = list(tk.iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == [73094]                                   # not the sushi mat
    assert rows[0][8].startswith("(no brand: frozen product) frozen")
    assert rows[0][7] == "[Vietnamese] nước mía đông lạnh\n[English] frozen sugarcane juice"
    rn = list(wb["RedNote Search List"].iter_rows(min_row=2, values_only=True))
    assert [(r[0], r[7]) for r in rn] == [(73094, "[Chinese] 冷冻甘蔗汁")]
    header = [c.value for c in wb[MAIN][1]]
    assert header[12:] == ["English 1", "English 2", "English 3", "Vietnamese 1", "Vietnamese 2",
                           "Vietnamese 3", "Chinese 1", "Chinese 2", "Chinese 3"]


def test_khmer_phrases_are_not_searched(data, tmp_path):
    """Khmer is detected but not searched: off the TikTok list, and no Khmer block."""
    data["items"] = [_item("OCT NEW", 1, "21033", "FISH SCE TREY KANCHANHCHRAS (B",
                           search_queries="ទឹកត្រី trey kanchanhchras|trey kanchanhchras fish sauce|"
                                          "Trey Kanchanhchras 鱼露")]
    data["rows"] = [_row("OCT NEW", 1, "21033", "FISH SCE TREY KANCHANHCHRAS (B", "Unmatched")]
    path = tmp_path / "khmer.xlsx"
    ex.build_workbook(data).save(path)
    wb = load_workbook(path)
    assert wb["TikTok Search List"]["H2"].value == "[English] trey kanchanhchras fish sauce"
    assert wb["RedNote Search List"]["H2"].value == "[Chinese] Trey Kanchanhchras 鱼露"
    assert "q=trey%20kanchanhchras%20fish%20sauce" in wb["TikTok Search List"]["L2"].value
    assert not any(str(c.value).startswith("Khmer") for c in wb[MAIN][1])
    assert ex.search_phrases("ទឹកត្រី|fish sauce") == ["fish sauce"]


def test_japanese_phrases_are_not_searched(data, tmp_path):
    """A Japanese import's old home-market phrase mixes kana and kanji. It is not searched
    and must not be filed under Chinese (NOV 2026: Hata, Morinaga Moonlight landed in the
    Chinese block)."""
    data["items"] = [_item("OCT NEW", 1, "34705", "HATA SODA WHITE PEACH", brand="Hata",
                           search_queries="ハタ鉱泉 ソーダ 白桃|hata white peach soda|Hata 白桃汽水")]
    data["rows"] = [_row("OCT NEW", 1, "34705", "HATA SODA WHITE PEACH", "Unmatched")]
    path = tmp_path / "japanese.xlsx"
    ex.build_workbook(data).save(path)
    wb = load_workbook(path)
    assert wb["TikTok Search List"]["H2"].value == "[English] hata white peach soda"
    assert wb["RedNote Search List"]["H2"].value == "[Chinese] Hata 白桃汽水"
    assert [c.value for c in wb[MAIN][1]][12::3] == ["English 1", "Chinese 1"]
    assert ex.search_phrases("ハタ鉱泉 ソーダ 白桃|Hata 白桃汽水") == ["Hata 白桃汽水"]


@pytest.mark.parametrize("phrase, language", [
    ("fanta grape", "English"), ("แฟนต้า องุ่น", "Thai"), ("芬达 葡萄", "Chinese"),
    ("kẹo dẻo hoshi xoài", "Vietnamese"), ("Bánh Tráng", "Vietnamese"), ("ទឹកត្រី", "Khmer"),
    ("Sunflower 沙丁鱼罐头", "Chinese"), ("keo deo hoshi", "English"),
    # kana makes it Japanese however much kanji it carries; kanji alone stays Chinese
    ("ハタ鉱泉 ソーダ グレープ", "Japanese"), ("ハタ鉱泉 ソーダ 白桃", "Japanese"),
    ("森永 ムーンライト さつまいも", "Japanese"), ("ｿｰﾀﾞ", "Japanese"),
    ("Hata 葡萄汽水", "Chinese"), ("乐事 烤牛肉味 薯片", "Chinese"), ("芬达 白桃 汽水", "Chinese"),
    ("森永 白桃", "Chinese"),
    ("ฮันนี่บี ชานมไข่มุก", "Thai"), ("nước mía đông lạnh", "Vietnamese"),
    ("marca piña toyo", "English")])                      # ñ is not a Vietnamese letter
def test_phrase_language(phrase, language):
    assert ex.phrase_language(phrase) == language
