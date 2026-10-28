"""Check a monthly promo INPUT file against the schema (promo_input_schema.py).

Run it on the merged workbook before the monthly job. ERRORS stop the load
(load_promo_list.py runs the same check); WARNINGS are worth a look but do not.

    python scripts/promo/validate_promo_input.py "C:/Users/me/Downloads/UST 2026 NOV PROMO INPUT.xlsx"

Also the reader: read_input() returns the rows, the meta and the photos, for the
loader. Photos come two ways and both are read: a picture PLACED IN the cell
(Excel "Place in Cell", stored as rich data, invisible to openpyxl) and a picture
floating over the cell (anchored on it).
"""

from __future__ import annotations

import os
import posixpath
import re
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path.cwd() / "scripts" / "promo"))
try:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
except NameError:
    pass

from promo_input_schema import (  # noqa: E402
    COLUMN_NAMES, COLUMNS, META_REQUIRED, PROMO_LISTS, REQUIRED, SCHEMA_VERSION, SHEET_ITEMS, SHEET_META,
    SHEET_UNMATCHED,
)

FILENAME = re.compile(r"^UST (20\d{2}) (JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC) PROMO INPUT\.xlsx$", re.I)
MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]


# ─────────────────────────────────────────────────────────────────────────────
# Read
# ─────────────────────────────────────────────────────────────────────────────

def _text(v) -> str | None:
    """A cell as text: dates as YYYY-MM-DD, whole floats without '.0', blanks as None."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def _sheet_xml_path(z: zipfile.ZipFile, title: str) -> str | None:
    wb = z.read("xl/workbook.xml").decode("utf-8")
    m = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="(rId\d+)"' % re.escape(title), wb)
    if not m:
        return None
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    t = re.search(r'Id="%s"[^>]*Target="([^"]+)"|Target="([^"]+)"[^>]*Id="%s"' % (m.group(1), m.group(1)), rels)
    target = (t.group(1) or t.group(2)).lstrip("/")
    return target if target.startswith("xl/") else posixpath.normpath(posixpath.join("xl", target))


def _in_cell_pictures(path: str, title: str, col_letter: str) -> dict[int, tuple[bytes, str]]:
    """{row: (bytes, ext)} for pictures placed IN the cells of one column (Excel rich
    data: cell vm= -> metadata valueMetadata -> futureMetadata rvb -> rich value ->
    richValueRel -> media file)."""
    out = {}
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        sheet = _sheet_xml_path(z, title)
        if not sheet or "xl/metadata.xml" not in names or "xl/richData/rdrichvalue.xml" not in names:
            return out
        xml = z.read(sheet).decode("utf-8")
        meta = z.read("xl/metadata.xml").decode("utf-8")
        value_meta = meta.split("<valueMetadata", 1)[1] if "<valueMetadata" in meta else ""
        vm_to_bk = re.findall(r'<bk><rc t="\d+" v="(\d+)"/></bk>', value_meta)
        rvb = re.findall(r'<xlrd:rvb i="(\d+)"/>', meta)
        rich = re.findall(r"<rv [^>]*>(.*?)</rv>", z.read("xl/richData/rdrichvalue.xml").decode("utf-8"), re.S)
        rel_ids = re.findall(r'r:id="(rId\d+)"', z.read("xl/richData/richValueRel.xml").decode("utf-8"))
        rels = z.read("xl/richData/_rels/richValueRel.xml.rels").decode("utf-8")
        targets = {}
        for rel in re.findall(r"<Relationship [^>]*/>", rels):
            rid, tgt = re.search(r'Id="([^"]+)"', rel), re.search(r'Target="([^"]+)"', rel)
            if rid and tgt:
                targets[rid.group(1)] = posixpath.normpath(posixpath.join("xl/richData", tgt.group(1)))
        for row, vm in re.findall(r'<c r="%s(\d+)"[^>]*\bvm="(\d+)"' % col_letter, xml):
            try:
                rv = rich[int(rvb[int(vm_to_bk[int(vm) - 1])])]
                rel = int(re.findall(r"<v[^>]*>(\d+)</v>", rv)[0])
                media = targets[rel_ids[rel]]
                out[int(row)] = (z.read(media), media.rsplit(".", 1)[-1].lower())
            except (IndexError, KeyError, ValueError):
                continue
    return out


def read_input(path: str) -> dict:
    """{rows: [{column: text, '_row': n}], header: [...], meta: {key: [values]},
    unmatched: n, photos: {row: (bytes, ext)}}. Never raises on content; validate()
    judges it."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.load_workbook(path, data_only=True)       # not read-only: anchored images
    out = {"sheets": wb.sheetnames, "rows": [], "header": [], "meta": {}, "unmatched": 0, "photos": {},
           "floating_photos": 0}
    if SHEET_META in wb.sheetnames:
        for r in wb[SHEET_META].iter_rows(min_row=2, values_only=True):
            if r and _text(r[0]):
                out["meta"].setdefault(_text(r[0]).lower(), []).append(_text(r[1]) if len(r) > 1 else None)
    if SHEET_UNMATCHED in wb.sheetnames:
        out["unmatched"] = sum(1 for r in wb[SHEET_UNMATCHED].iter_rows(min_row=2, values_only=True)
                               if r and any(v not in (None, "") for v in r))
    if SHEET_ITEMS not in wb.sheetnames:
        return out
    ws = wb[SHEET_ITEMS]
    header = [(_text(c.value) or "").lower() for c in ws[1]]
    out["header"] = header
    for row in ws.iter_rows(min_row=2):
        vals = {h: _text(c.value) for h, c in zip(header, row) if h}
        if not any(v for k, v in vals.items() if k != "photo"):
            continue                                        # blank trailing row
        vals["_row"] = row[0].row
        out["rows"].append(vals)
    if "photo" in header:
        col = header.index("photo")
        letter = get_column_letter(col + 1)
        photos = _in_cell_pictures(path, ws.title, letter)
        floating = 0
        for img in getattr(ws, "_images", []):              # pictures floating on the cell
            frm = img.anchor._from
            if frm.col == col and frm.row + 1 not in photos:
                data = img._data()
                photos[frm.row + 1] = (data, (img.format or "png").lower())
                floating += 1
        out["photos"] = photos
        out["floating_photos"] = floating
    return out


# ─────────────────────────────────────────────────────────────────────────────
# Validate
# ─────────────────────────────────────────────────────────────────────────────

def _is_date(v: str) -> bool:
    try:
        datetime.strptime(v[:10], "%Y-%m-%d")
        return len(v) == 10 or v[10:].strip() in ("", "00:00:00")
    except ValueError:
        return False


def validate(path: str, data: dict | None = None) -> tuple[list[str], list[str]]:
    """-> (errors, warnings)."""
    data = data or read_input(path)
    errors, warnings = [], []
    name = os.path.basename(path)

    for sheet in (SHEET_ITEMS, SHEET_META):
        if sheet not in data["sheets"]:
            errors.append(f"no sheet '{sheet}' (sheets: {data['sheets']})")
    if errors:
        return errors, warnings

    # meta
    meta = data["meta"]
    for key in META_REQUIRED:
        if not any(meta.get(key) or []):
            errors.append(f"meta: no '{key}' row")
    version = (meta.get("schema_version") or [None])[0]
    if version and version != SCHEMA_VERSION:
        errors.append(f"meta: schema_version {version}, this pipeline reads {SCHEMA_VERSION}")
    month = (meta.get("promo_month") or [None])[0]
    if month and not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", month):
        errors.append(f"meta: promo_month '{month}' is not YYYY-MM")

    # file name, when it follows the convention, must agree with the month
    m = FILENAME.match(name)
    if not m:
        warnings.append(f"file name '{name}' is not 'UST <YYYY> <MON> PROMO INPUT.xlsx'")
    elif month and f"{m.group(1)}-{MONTHS.index(m.group(2).upper()) + 1:02d}" != month:
        errors.append(f"file name says {m.group(1)} {m.group(2).upper()}, meta promo_month says {month}")

    # columns
    header = data["header"]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        errors.append(f"items: missing required columns {missing}")
    absent = [c for c in COLUMN_NAMES if c not in header and c not in REQUIRED]
    if absent:
        warnings.append(f"items: optional columns not present (fine if no source had them): {absent}")
    extra = [h for h in header if h and h not in COLUMN_NAMES]
    if extra:
        warnings.append(f"items: columns not in the schema are ignored: {extra}")
    dupes = sorted({h for h in header if h and header.count(h) > 1})
    if dupes:
        errors.append(f"items: columns appear twice: {dupes}")
    if missing or dupes:
        return errors, warnings

    rows = data["rows"]
    if not rows:
        errors.append("items: no rows")
        return errors, warnings

    # rows
    kinds = {c.name: c.kind for c in COLUMNS}
    seen = {}
    for r in rows:
        at = f"items row {r['_row']}"
        for col in REQUIRED:
            if not r.get(col):
                errors.append(f"{at}: {col} is blank")
        for col, v in r.items():
            if col.startswith("_") or v is None or col not in kinds:
                continue
            kind = kinds[col]
            if kind == "item_no" and not v.isdigit():
                errors.append(f"{at}: item_no '{v}' is not digits only")
            elif kind == "month" and month and v != month:
                errors.append(f"{at}: promo_month '{v}' differs from meta ({month})")
            elif kind == "month" and not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", v):
                errors.append(f"{at}: promo_month '{v}' is not YYYY-MM")
            elif kind == "list" and v.upper() not in PROMO_LISTS:
                errors.append(f"{at}: promo_list '{v}' is not one of {list(PROMO_LISTS)}")
            elif kind == "date" and not _is_date(v):
                errors.append(f"{at}: {col} '{v}' is not a date (YYYY-MM-DD)")
            elif kind == "code" and not re.fullmatch(r"[A-Za-z]{2}", v):
                warnings.append(f"{at}: {col} '{v}' is not a 2-letter country code")
            elif kind == "pipe_list" and ("\n" in v or (";" in v and "|" not in v)):
                warnings.append(f"{at}: {col} should be joined with ' | ', got '{v[:60]}'")
        key = ((r.get("promo_list") or "").upper(), r.get("item_no"))
        if all(key):
            if key in seen:
                errors.append(f"{at}: {key[0]} item {key[1]} is already on row {seen[key]}")
            seen.setdefault(key, r["_row"])

    # coverage, so a thin month is visible before the run
    photos = data["photos"]
    detail = sum(1 for r in rows if r.get("short_description") or r.get("selling_points") or r.get("notes"))
    lists = {}
    for r in rows:
        lists[(r.get("promo_list") or "?").upper()] = lists.get((r.get("promo_list") or "?").upper(), 0) + 1
    warnings.append(f"coverage: {len(rows)} rows ({', '.join(f'{k} {v}' for k, v in lists.items())}); "
                    f"{detail} with a description or selling points, {len(photos)} with a photo")
    if data.get("floating_photos"):
        warnings.append(f"{data['floating_photos']} photo(s) float over their cells instead of sitting in them: "
                        f"they load, but are easy to miss and do not move with their rows. Run "
                        f"python scripts/promo/photos_in_cells.py \"{path}\"")
    if data["unmatched"]:
        warnings.append(f"{SHEET_UNMATCHED}: {data['unmatched']} description-file row(s) are on no promo "
                        f"list and will NOT be loaded")
    return errors, warnings


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("path")
    args = ap.parse_args(argv)
    errors, warnings = validate(args.path)
    for w in warnings:
        print(f"  warning: {w}")
    for e in errors:
        print(f"  ERROR: {e}")
    print(f"{os.path.basename(args.path)}: " + (f"{len(errors)} error(s): NOT ready to load" if errors
                                                 else "valid, ready to load"))
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
