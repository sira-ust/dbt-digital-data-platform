"""Put an INPUT workbook's photos IN their cells (Excel "Place in Cell").

A merge made with openpyxl (as Claude does with prompts/promo_merge_input.md) can only
FLOAT a picture over its cell. The pipeline reads either kind, but a floating picture on
the 17th column of a 15-point row is easy to miss, and it does not move with its row
when the sheet is sorted or filtered. This rewrites the file so each photo sits inside
its `photo` cell, the way the description file stores them, and makes the rows tall
enough to see it.

    python scripts/promo/photos_in_cells.py "C:/Users/me/Desktop/UST 2026 NOV PROMO INPUT.xlsx"

The file is rewritten in place (pass --out to write elsewhere). Image bytes are copied
unchanged; only the photo column of the items sheet is touched. Excel stores an in-cell
picture as rich data: the cell holds an error value with a value-metadata index (vm=),
xl/metadata.xml points that index at a rich value, and the rich value points at the
image through xl/richData/richValueRel.xml. The fixed parts below are copied from a
workbook Excel saved (the NOV 2026 description file).
"""

from __future__ import annotations

import argparse
import io
import os
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path.cwd() / "scripts" / "promo"))
try:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
except NameError:
    pass

from promo_input_schema import SHEET_ITEMS  # noqa: E402

ROW_HEIGHT, COL_WIDTH = 80, 16          # points / characters: room for a visible picture

RV_STRUCTURE = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<rvStructures xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" count="1">'
                '<s t="_localImage"><k n="_rvRel:LocalImageIdentifier" t="i"/><k n="CalcOrigin" t="i"/></s>'
                '</rvStructures>')
RV_TYPES = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<rvTypesInfo xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata2" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="x" '
            'xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><global><keyFlags>'
            + "".join(f'<key name="{k}">' + ('<flag name="ExcludeFromFile" value="1"/>' if k == "_Self" else "")
                      + '<flag name="ExcludeFromCalcComparison" value="1"/></key>'
                      for k in ("_Self", "_DisplayString", "_Flags", "_Format", "_SubLabel", "_Attribution",
                                "_Icon", "_Display", "_CanonicalPropertyNames", "_ClassificationId"))
            + "</keyFlags></global></rvTypesInfo>")
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WB_RELS = [  # (type, target)
    (f"{REL}/sheetMetadata", "metadata.xml"),
    ("http://schemas.microsoft.com/office/2022/10/relationships/richValueRel", "richData/richValueRel.xml"),
    ("http://schemas.microsoft.com/office/2017/06/relationships/rdRichValue", "richData/rdrichvalue.xml"),
    ("http://schemas.microsoft.com/office/2017/06/relationships/rdRichValueStructure", "richData/rdrichvaluestructure.xml"),
    ("http://schemas.microsoft.com/office/2017/06/relationships/rdRichValueTypes", "richData/rdRichValueTypes.xml"),
]
OVERRIDES = {
    "/xl/metadata.xml": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml",
    "/xl/richData/richValueRel.xml": "application/vnd.ms-excel.richvaluerel+xml",
    "/xl/richData/rdrichvalue.xml": "application/vnd.ms-excel.rdrichvalue+xml",
    "/xl/richData/rdrichvaluestructure.xml": "application/vnd.ms-excel.rdrichvaluestructure+xml",
    "/xl/richData/rdRichValueTypes.xml": "application/vnd.ms-excel.rdrichvaluetypes+xml",
}
IMAGE_TYPES = {"png": "image/png", "jpeg": "image/jpeg", "jpg": "image/jpeg", "gif": "image/gif"}


def _metadata(n: int) -> str:
    future = "".join('<bk><extLst><ext uri="{3e2802c4-a4d2-4d8b-9148-e3be6c30e623}">'
                     f'<xlrd:rvb i="{i}"/></ext></extLst></bk>' for i in range(n))
    value = "".join(f'<bk><rc t="1" v="{i}"/></bk>' for i in range(n))
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata">'
            '<metadataTypes count="1"><metadataType name="XLRICHVALUE" minSupportedVersion="120000" copy="1" '
            'pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" '
            'clearComments="1" assign="1" coerce="1"/></metadataTypes>'
            f'<futureMetadata name="XLRICHVALUE" count="{n}">{future}</futureMetadata>'
            f'<valueMetadata count="{n}">{value}</valueMetadata></metadata>')


def place_photos_in_cells(src: str, out: str | None = None) -> int:
    """Rewrite src (or write out) with the items sheet's floating photos placed in
    their photo cells. -> the number of photos placed. A file with no floating photos
    on the photo column is left as it is (0)."""
    import openpyxl
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    out = out or src
    wb = openpyxl.load_workbook(src)
    ws = wb[SHEET_ITEMS]
    header = [str(c.value or "").strip().lower() for c in ws[1]]
    if "photo" not in header:
        raise SystemExit(f"{os.path.basename(src)}: the items sheet has no 'photo' column")
    col = header.index("photo")
    letter = get_column_letter(col + 1)
    photos, keep = {}, []
    for img in ws._images:
        frm = img.anchor._from
        if frm.col == col and frm.row >= 1 and frm.row + 1 not in photos:
            photos[frm.row + 1] = (img._data(), (img.format or "png").lower())
        else:
            keep.append(img)
    if not photos:
        return 0
    ws._images = keep
    for row in photos:                       # a styled cell is always written, so it can be found
        cell = ws[f"{letter}{row}"]
        cell.value = None
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = ROW_HEIGHT
    ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, COL_WIDTH)

    with tempfile.TemporaryDirectory() as tmp:
        plain = Path(tmp) / "plain.xlsx"
        wb.save(plain)
        sheet_path = _items_sheet_path(plain)
        rows = sorted(photos)
        with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(Path(tmp) / "out.xlsx", "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == sheet_path:
                    data = _mark_cells(data.decode("utf-8"), letter, rows).encode("utf-8")
                elif item.filename == "xl/_rels/workbook.xml.rels":
                    data = _add_wb_rels(data.decode("utf-8")).encode("utf-8")
                elif item.filename == "[Content_Types].xml":
                    data = _add_types(data.decode("utf-8"), {photos[r][1] for r in rows}).encode("utf-8")
                zout.writestr(item, data)
            rels, rvs = [], []
            for i, row in enumerate(rows):
                data, ext = photos[row]
                name = f"rv_image{i + 1}.{ext}"
                zout.writestr(f"xl/media/{name}", data)
                rels.append(f'<Relationship Id="rId{i + 1}" Type="{REL}/image" Target="../media/{name}"/>')
                rvs.append(f'<rv s="0"><v>{i}</v><v>5</v></rv>')
            n = len(rows)
            zout.writestr("xl/metadata.xml", _metadata(n))
            zout.writestr("xl/richData/rdrichvalue.xml",
                          '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<rvData xmlns='
                          f'"http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" count="{n}">'
                          + "".join(rvs) + "</rvData>")
            zout.writestr("xl/richData/rdrichvaluestructure.xml", RV_STRUCTURE)
            zout.writestr("xl/richData/rdRichValueTypes.xml", RV_TYPES)
            zout.writestr("xl/richData/richValueRel.xml",
                          '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<richValueRels xmlns='
                          '"http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel" '
                          f'xmlns:r="{REL}">' + "".join(f'<rel r:id="rId{i + 1}"/>' for i in range(n))
                          + "</richValueRels>")
            zout.writestr("xl/richData/_rels/richValueRel.xml.rels",
                          '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns='
                          '"http://schemas.openxmlformats.org/package/2006/relationships">' + "".join(rels)
                          + "</Relationships>")
        shutil.copyfile(Path(tmp) / "out.xlsx", out)
    return len(photos)


def _items_sheet_path(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        wb = z.read("xl/workbook.xml").decode("utf-8")
        rid = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="(rId\d+)"' % SHEET_ITEMS, wb).group(1)
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        rel = next(r for r in re.findall(r"<Relationship [^>]*/>", rels) if f'Id="{rid}"' in r)
        target = re.search(r'Target="([^"]+)"', rel).group(1).lstrip("/")
        return target if target.startswith("xl/") else f"xl/{target}"


def _mark_cells(xml: str, letter: str, rows: list[int]) -> str:
    """Each photo cell becomes <c r=".." s=".." t="e" vm="k"><v>#VALUE!</v></c>."""
    for k, row in enumerate(rows, start=1):
        ref = f"{letter}{row}"
        pat = re.compile(r'<c r="%s"([^>]*?)(?:/>|>.*?</c>)' % ref, re.S)
        m = pat.search(xml)
        if not m:
            raise SystemExit(f"photo cell {ref} was not written; cannot place its picture")
        style = re.search(r'\ss="\d+"', m.group(1))
        xml = xml[:m.start()] + f'<c r="{ref}"{style.group(0) if style else ""} t="e" vm="{k}"><v>#VALUE!</v></c>' \
            + xml[m.end():]
    return xml


def _add_wb_rels(xml: str) -> str:
    used = {int(i) for i in re.findall(r'Id="rId(\d+)"', xml)}
    nxt = max(used or {0}) + 1
    add = ""
    for typ, target in WB_RELS:
        if f'Target="{target}"' not in xml:
            add += f'<Relationship Id="rId{nxt}" Type="{typ}" Target="{target}"/>'
            nxt += 1
    return xml.replace("</Relationships>", add + "</Relationships>")


def _add_types(xml: str, exts: set[str]) -> str:
    add = ""
    for ext in sorted(exts):
        if f'Extension="{ext}"' not in xml:
            add += f'<Default Extension="{ext}" ContentType="{IMAGE_TYPES.get(ext, "image/" + ext)}"/>'
    for part, ctype in OVERRIDES.items():
        if f'PartName="{part}"' not in xml:
            add += f'<Override PartName="{part}" ContentType="{ctype}"/>'
    return xml.replace("</Types>", add + "</Types>")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("path")
    ap.add_argument("--out", help="write here instead of rewriting the file in place")
    args = ap.parse_args(argv)
    n = place_photos_in_cells(args.path, args.out)
    print(f"{os.path.basename(args.out or args.path)}: "
          + (f"{n} photo(s) placed in their cells" if n else "no floating photos on the photo column; unchanged"))


if __name__ == "__main__":
    main()
