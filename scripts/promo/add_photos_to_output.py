"""Add the product photos to a finished output workbook, from the month's INPUT file.

The export puts each item's photo beside its name (the Photo column) since
2026-10-28. A workbook exported before that has none; this adds them, matched by
ITEM NUMBER (never by description: 35060 and 35061 are both HOSHI GUMMY MANGO, each
with its own photo):

    python scripts/promo/add_photos_to_output.py "UST_NOV_Promo_Social_Listening_Match.xlsx" "UST 2026 NOV PROMO INPUT.xlsx"

- The workbook has a Photo column: the photos go in it.
- It has none: a Photo column is added AFTER the last column, so no formula, link or
  note on the sheet moves.

Photos float over their cells (sized to the cell, on a taller row), the kind openpyxl
keeps when it saves, so a later edit by Claude does not drop them. An item that already
has a photo on its cell is left as it is: safe to run twice. The file is rewritten in
place unless --out is given.
"""

from __future__ import annotations

import argparse
import io
import os
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path.cwd() / "scripts" / "promo"))
try:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
except NameError:
    pass

from validate_promo_input import read_input  # noqa: E402

SHEET_MAIN = "Promo vs Social Listening"      # export_promo_social_match.SHEET_MAIN
PHOTO_PX, PHOTO_ROW_HEIGHT, PHOTO_COL_WIDTH = 76, 60, 14


def photos_by_item(input_path: str) -> dict[str, bytes]:
    """{item_no: photo bytes} from an INPUT file (in-cell or floating photos)."""
    data = read_input(input_path)
    row_item = {r["_row"]: str(r["item_no"]) for r in data["rows"] if r.get("item_no")}
    return {row_item[row]: blob for row, (blob, _) in data["photos"].items() if row in row_item}


def add_photos(result_path: str, input_path: str, out: str | None = None) -> tuple[int, str]:
    """-> (photos added, the Photo column's letter)."""
    import openpyxl
    from copy import copy
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter

    photos = photos_by_item(input_path)
    wb = openpyxl.load_workbook(result_path)
    if SHEET_MAIN not in wb.sheetnames:
        raise SystemExit(f"{os.path.basename(result_path)}: no sheet '{SHEET_MAIN}'")
    ws = wb[SHEET_MAIN]
    header = [str(c.value or "").strip() for c in ws[1]]
    if "Item No." not in header:
        raise SystemExit(f"{os.path.basename(result_path)}: no 'Item No.' column on {SHEET_MAIN}")
    item_col = header.index("Item No.") + 1
    if "Photo" in header:
        col = header.index("Photo") + 1
    else:                                    # after the last column: nothing on the sheet moves
        col = ws.max_column + 1
        cell, left = ws.cell(row=1, column=col, value="Photo"), ws.cell(row=1, column=col - 1)
        cell.font, cell.fill, cell.border, cell.alignment = (copy(left.font), copy(left.fill),
                                                              copy(left.border), copy(left.alignment))
        ws.column_dimensions[get_column_letter(col)].width = PHOTO_COL_WIDTH
    letter = get_column_letter(col)
    taken = {(i.anchor._from.col, i.anchor._from.row) for i in ws._images}
    added = 0
    for r in range(2, ws.max_row + 1):
        item = ws.cell(row=r, column=item_col).value
        blob = photos.get(str(item).strip().removesuffix(".0")) if item is not None else None
        if not blob or (col - 1, r - 1) in taken:
            continue
        img = XLImage(io.BytesIO(blob))
        scale = PHOTO_PX / max(img.height, 1)
        img.width, img.height = int(img.width * scale), PHOTO_PX
        ws.add_image(img, f"{letter}{r}")
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, PHOTO_ROW_HEIGHT)
        added += 1
    with tempfile.TemporaryDirectory() as tmp:
        tmp_out = Path(tmp) / "out.xlsx"
        wb.save(tmp_out)
        shutil.copyfile(tmp_out, out or result_path)
    return added, letter


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("result", help="the output workbook (UST_<MON>_Promo_Social_Listening_Match.xlsx)")
    ap.add_argument("input", help="the month's INPUT file, with the photos")
    ap.add_argument("--out", help="write here instead of rewriting the result in place")
    args = ap.parse_args(argv)
    n, letter = add_photos(args.result, args.input, args.out)
    print(f"{os.path.basename(args.out or args.result)}: {n} photo(s) added in column {letter}")


if __name__ == "__main__":
    main()
