"""The workbook export: layout, formatting and live formulas, from synthetic rows.

No database needed — build_workbook() takes the same dict read_inputs() returns.
The formula test needs the optional `formulas` package (a spreadsheet engine) and
is skipped without it. Run:  python -m pytest scripts/tests
"""

import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import export_promo_social_match as ex  # noqa: E402


def _row(sheet, row_order, item_no, desc, status, rank=1, **match):
    base = {"promo_sheet": sheet, "row_order": row_order, "match_rank": rank, "item_no": item_no,
            "item_no_int": int(item_no), "promo_description": desc + "  ", "promo_size": "24X10.9OZ",
            "promo_code": None if sheet.endswith("NEW") else "5+1", "match_status": status,
            "matched_content_item": None, "channel": None, "profile": None, "post_date": None,
            "engagement_text": None, "caption": None, "link": None, "source": None,
            "match_note": None, "views": None, "likes": None}
    base.update(match)
    # the mart's item-level columns; one source per item unless a test overrides them
    src = base["source"]
    base.setdefault("item_match_status", status)
    base.setdefault("social_match_status", status if src == "Social Listening" else None)
    base.setdefault("scraped_match_status", status if src == "Scraped" else None)
    return base


def _item(sheet, row_order, item_no, desc, branded=True, **kw):
    base = {"promo_sheet": sheet, "row_order": row_order, "item_no": item_no, "promo_description": desc,
            "promo_size": "24X10.9OZ", "is_branded": branded, "brand": "Fanta" if branded else None,
            "product": "soft drink", "flavor_variant": "grape", "search_queries": "แฟนต้า องุ่น|fanta grape",
            "brand_terms": "fanta|แฟนต้า", "flavor_terms": "grape|องุ่น", "negative_terms": None,
            "source_file": "UST 2026 OCT NEW PROMO.xlsx"}
    base.update(kw)
    return base


@pytest.fixture
def data():
    sl = dict(matched_content_item="ดีโด้ น้ำส้มสายน้ำผึ้ง (Deedo Sainamphueng orange juice)",
              channel="Web", profile="trueid.net", post_date="2026-07-13", engagement_text="Not reported",
              caption="มิกซ์ กาแฟส้ม", link="https://www.trueid.net/watch/th-th/short/8yXZnXjx325b",
              source="Social Listening", match_note="Brand, product and flavor match.", views=0, likes=0)
    yt = dict(matched_content_item="Fanta Grape", channel="YouTube", profile="Fanta TH", post_date="2022-03-14",
              engagement_text="176,000 views, 900 likes", caption="ระเบิดความซ่ากับองุ่นป๊อบ!",
              link="https://www.youtube.com/watch?v=GUwN8yYJboI", source="Scraped",
              match_note="Brand + product/flavor named in video.", views=176000, likes=900)
    rows = [
        _row("OCT NEW", 1, "86251", "BAMBOO SUSHI MAT", "Unmatched"),
        _row("OCT NEW", 2, "37353", "BENTO SQUID SNACK (BBQ CHEESE)", "Close",
             **{**sl, "matched_content_item": "Mực Bento", "match_note": "Same brand/product line"}),
        _row("OCT PROMO", 1, "34585", "DEEDO SAINAMPHUENG FRUIT JUICE", "Exact", **sl),
        _row("OCT PROMO", 2, "34662", "FANTA GRAPE FLAVOR", "Exact", 1, **yt),
        _row("OCT PROMO", 2, "34662", "FANTA GRAPE FLAVOR", "Exact", 2, **{**yt, "link": "https://youtu.be/x"}),
        _row("OCT PROMO", 3, "34658", "FANTA ORANGE FLAVOR", "Unmatched"),
    ]
    items = [_item("OCT NEW", 1, "86251", "BAMBOO SUSHI MAT", branded=False),
             _item("OCT NEW", 2, "37353", "BENTO SQUID SNACK (BBQ CHEESE)"),
             _item("OCT PROMO", 1, "34585", "DEEDO SAINAMPHUENG FRUIT JUICE"),
             _item("OCT PROMO", 2, "34662", "FANTA GRAPE FLAVOR"),
             _item("OCT PROMO", 3, "34658", "FANTA ORANGE FLAVOR", flavor_terms=None)]
    return {"promo_month": "2026-10-01", "rows": rows, "items": items,
            "mentions": {"n": 53601, "first_at": datetime(2026, 2, 20), "last_at": datetime(2026, 9, 29)},
            "searched_at": datetime(2026, 10, 2, 5, 0)}


@pytest.fixture
def saved(data, tmp_path):
    path = tmp_path / ex.output_name(data["promo_month"])
    ex.build_workbook(data).save(path)
    return path


def test_file_name(data):
    assert ex.output_name(data["promo_month"]) == "UST_OCT_Promo_Social_Listening_Match.xlsx"


def test_sheets_and_header(saved):
    wb = load_workbook(saved)
    assert wb.sheetnames == ["Summary", "Promo vs Social Listening", "TikTok Search List"]
    assert wb.active.title == "Summary"                                   # opens on the Summary
    ws = wb["Promo vs Social Listening"]
    assert [c.value for c in ws[1]] == [
        "Promo Sheet", "Item No.", "OCT Promo Item", "Size", "Promo Code", "Match Status",
        "On Trending Social (Mentionlytics)", "Direct Scraping (TikTok & YouTube)",
        "Matched Content Item", "Channel", "Profile", "Post Date", "Views", "Likes",
        "Approx. Count", "Engagement (as shown)", "Caption", "Link", "Source", "Match Note"]
    h = ws["A1"]
    assert h.font.bold and h.font.name == "Arial" and h.font.color.rgb == "FFFFFFFF"
    assert h.fill.fgColor.rgb == "FF1F4E78" and h.alignment.vertical == "center" and h.alignment.wrap_text
    assert ws.row_dimensions[1].height == 27.75


def test_layout(saved):
    ws = load_workbook(saved)["Promo vs Social Listening"]
    widths = [ws.column_dimensions[c].width for c in "ABCDEFGHIJKLMNOPQRST"]
    assert widths == [13, 10, 32, 15, 10, 12, 16, 16, 36, 14, 22, 11, 13, 10, 9, 22, 55, 11, 18, 48]
    assert ws.freeze_panes == "D2"
    assert ws.auto_filter.ref == f"A1:T{ws.max_row}"
    assert all(ws.row_dimensions[r].height == 39.75 for r in range(2, ws.max_row + 1))


def test_rows_order_types_and_formatting(saved):
    ws = load_workbook(saved)["Promo vs Social Listening"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["OCT NEW"] * 2 + ["OCT PROMO"] * 4      # NEW first, source order
    assert all(isinstance(r[1], int) for r in rows)                        # Item No. is an integer
    assert rows[0][2] == "BAMBOO SUSHI MAT"                                 # trimmed
    assert rows[0][4] is None                                               # blank promo code
    assert all(v is None for v in rows[0][6:])                              # Unmatched: blanks
    assert rows[3][1] == rows[4][1] == 34662                                # consecutive rows per item
    body = ws["C3"]
    assert body.font.name == "Arial" and body.font.size == 10
    assert body.alignment.vertical == "top" and body.alignment.wrap_text
    assert body.border.left.color.rgb == "FFBFBFBF"
    assert ws["C2"].font.color.rgb == "FF808080"                            # Unmatched row greyed


def test_views_likes_and_engagement(saved):
    ws = load_workbook(saved)["Promo vs Social Listening"]
    yt = ws[5]                                                               # Fanta Grape, YouTube
    assert (yt[12].value, yt[13].value) == (176000, 900)
    assert yt[12].number_format == "#,##0" and yt[12].alignment.horizontal == "right"
    assert yt[14].value is None                                              # Approx. Count: only TikTok
    assert yt[15].value == "176,000 views, 900 likes"
    web = ws[4]                                                              # Deedo, a web article
    assert (web[12].value, web[13].value, web[15].value) == (None, None, "Not reported")


def test_banding_by_item(saved):
    ws = load_workbook(saved)["Promo vs Social Listening"]
    shaded = [ws[f"C{r}"].fill.fgColor.rgb == "FFEEF3F8" for r in range(2, ws.max_row + 1)]
    # items: sushi mat | bento | deedo | fanta grape x2 | fanta orange
    assert shaded == [False, True, False, True, True, False]


def test_status_styles_and_links(saved):
    ws = load_workbook(saved)["Promo vs Social Listening"]
    styles = {ws[f"F{r}"].value: (ws[f"F{r}"].fill.fgColor.rgb, ws[f"F{r}"].font.color.rgb)
              for r in range(2, ws.max_row + 1)}
    assert styles == {"Unmatched": ("FFEDEDED", "FF595959"), "Close": ("FFFFEB9C", "FF9C5700"),
                      "Exact": ("FFC6EFCE", "FF006100")}
    assert all(ws[f"F{r}"].font.bold for r in range(2, ws.max_row + 1))
    link = ws["R4"]
    assert link.value == "Open ↗" and link.hyperlink.target.startswith("https://www.trueid.net/")
    assert link.font.underline == "single" and link.font.color.rgb == "FF0563C1"
    assert ws["R2"].value is None and ws["R2"].hyperlink is None              # Unmatched: no link


def test_source_columns(saved):
    ws = load_workbook(saved)["Promo vs Social Listening"]
    # Deedo (row 4): social listening only; Fanta Grape (row 5): YouTube only
    assert (ws["G4"].value, ws["H4"].value, ws["S4"].value) == ("Exact", None, "On Trending Social")
    assert (ws["G5"].value, ws["H5"].value, ws["S5"].value) == (None, "Exact", "Direct Scraping")
    assert ws["G4"].fill.fgColor.rgb == "FFC6EFCE" and ws["G4"].font.bold
    assert (ws["G2"].value, ws["H2"].value) == (None, None)                  # Unmatched: both blank


def test_summary_formulas_and_footnote(saved):
    sm = load_workbook(saved)["Summary"]
    assert [c.value for c in sm[1]][:3] == ["Match Status", "Promo Items", "Matched Posts (rows)"]
    assert [sm[f"A{r}"].value for r in range(2, 6)] == ["Exact", "Close", "Unmatched", "Total"]
    assert sm["B2"].value.startswith("=SUMPRODUCT(") and "COUNTIFS(" in sm["B2"].value
    assert sm["C2"].value == "=IF(A2=\"Unmatched\",0,COUNTIF('Promo vs Social Listening'!F:F,A2))"
    assert sm["B5"].value == "=SUM(B2:B4)" and sm["C5"].value == "=SUM(C2:C4)"
    note = sm["A7"].value
    assert "UST 2026 OCT NEW PROMO.xlsx (5 items)" in note and "53,601 posts" in note
    assert "Oct 2, 2026" in note and "YouTube shows views; TikTok shows likes" in note


def test_summary_sections(saved):
    sm = load_workbook(saved)["Summary"]
    col_a = [sm.cell(r, 1).value for r in range(1, sm.max_row + 1)]
    assert sm["A9"].value == "Promo items by sheet"
    assert [c.value for c in sm[10]][:4] == ["Match Status", "OCT NEW", "OCT PROMO", "Total"]
    assert '"OCT NEW"' in sm["B11"].value and '"OCT PROMO"' in sm["C11"].value
    assert sm["D11"].value == "=SUM(B11:C11)" and sm["A14"].value == "Total"
    assert sm["A16"].value == "Matched posts by channel"
    # channels in the rows, plus TikTok (the TikTok step adds it), alphabetical
    assert col_a[17:21] == ["TikTok", "Web", "YouTube", "Total"]
    assert sm["B18"].value == ("=COUNTIFS('Promo vs Social Listening'!$J:$J,$A18,"
                               "'Promo vs Social Listening'!$F:$F,\"Exact\")")
    i = col_a.index("Unmatched promo items (no post found)") + 1
    assert [c.value for c in sm[i + 1]][:3] == ["Item No.", "OCT Promo Item", "Promo Sheet"]
    assert [sm.cell(i + 2, 1).value, sm.cell(i + 3, 1).value] == [86251, 34658]
    assert sm.cell(i + 4, 1).value.startswith("List captured when the workbook was built")


def test_tiktok_list_has_unmatched_branded_items_only(saved):
    tk = load_workbook(saved)["TikTok Search List"]
    items = [r[0] for r in tk.iter_rows(min_row=2, values_only=True)]
    # unbranded sushi mat: no; Bento + Deedo matched in social listening: no;
    # Fanta Grape (YouTube Exact) and Fanta Orange (Unmatched): yes
    assert items == [34662, 34658]
    assert tk["G2"].value == "Exact (YouTube)" and tk["G3"].value == "Unmatched"
    assert tk["J3"].value == "(no flavor to require)"
    assert tk["L2"].hyperlink.target.startswith("https://www.tiktok.com/search/video?q=")


def test_summary_formulas_recalculate(saved):
    formulas = pytest.importorskip("formulas")
    model = formulas.ExcelModel().loads(str(saved)).finish()
    sol = model.calculate()

    def val(ref):
        key = next(k for k in sol if k.upper().endswith(f"SUMMARY'!{ref}"))
        v = sol[key].value
        return float(v[0][0] if hasattr(v, "__getitem__") else v)

    # Exact: Deedo + Fanta Grape = 2 items, 3 rows; Close: Bento 1/1; Unmatched: 2 items, 0 rows
    assert [val(f"B{r}") for r in (2, 3, 4, 5)] == [2, 1, 2, 5]
    assert [val(f"C{r}") for r in (2, 3, 4, 5)] == [3, 1, 0, 4]
    # by sheet: OCT NEW has the sushi mat (Unmatched) and Bento (Close); OCT PROMO the rest
    assert [val(f"B{r}") for r in (11, 12, 13)] == [0, 1, 1]
    assert [val(f"C{r}") for r in (11, 12, 13)] == [2, 0, 1]
    # by channel: Web = Bento (Close) + Deedo (Exact); YouTube = 2 Fanta Grape rows
    assert [val("B19"), val("C19"), val("B20")] == [1, 1, 2]


def test_own_brand_shows_both_sources(data, tmp_path):
    """An own-brand item keeps its social listening post AND gets scraped: one item
    status (the better one) on every row, each source's result in its own column,
    and it stays on the TikTok list although social listening matched it."""
    sl = dict(matched_content_item="Bánh tráng Chuồn Chuồn", channel="TikTok", profile="@a",
              post_date="2026-09-01", engagement_text="12 likes", caption="chuồn chuồn",
              link="https://www.tiktok.com/@a/video/1", source="Social Listening",
              match_note="Same brand/product line", views=0, likes=12)
    yt = dict(matched_content_item="Chuồn Chuồn rice paper", channel="YouTube", profile="Chuon Chuon Eats",
              post_date="2026-06-03", engagement_text="505 views", caption="bánh tráng",
              link="https://www.youtube.com/watch?v=y", source="Scraped",
              match_note="Brand + product/flavor named in video.", views=505, likes=5)
    both = dict(item_match_status="Exact", social_match_status="Close", scraped_match_status="Exact")
    data["rows"] = [_row("OCT NEW", 1, "43220", "CHUON CHUON RICE PAPER", "Close", 1, **sl, **both),
                    _row("OCT NEW", 1, "43220", "CHUON CHUON RICE PAPER", "Exact", 2, **yt, **both)]
    data["items"] = [_item("OCT NEW", 1, "43220", "CHUON CHUON RICE PAPER", brand="Chuon Chuon",
                           is_own_brand=True)]
    path = tmp_path / "own.xlsx"
    ex.build_workbook(data).save(path)
    wb = load_workbook(path)
    ws = wb["Promo vs Social Listening"]
    assert [ws[f"F{r}"].value for r in (2, 3)] == ["Exact", "Exact"]
    assert [ws[f"G{r}"].value for r in (2, 3)] == ["Close", "Close"]
    assert [ws[f"H{r}"].value for r in (2, 3)] == ["Exact", "Exact"]
    assert [ws[f"S{r}"].value for r in (2, 3)] == ["On Trending Social", "Direct Scraping"]
    tk = wb["TikTok Search List"]
    assert tk["A2"].value == 43220
    assert tk["G2"].value == "Close (On Trending Social); Exact (YouTube)"
